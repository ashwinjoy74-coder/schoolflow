"""SchoolFlow editor for Excel timetables.

Class sheets are shown as weekly grids. Double-click a lesson to change it or
fill an empty slot. Teacher conflicts are checked and matching teacher sheets
and Requirements counts are updated together.
"""
import os
import re
try:  # The web server only uses the lesson rules, never the Tk editor window.
    import tkinter as tk
    from tkinter import messagebox, ttk
except ImportError:
    tk = messagebox = ttk = None
from openpyxl import load_workbook
from openpyxl.cell.cell import MergedCell
from openpyxl.utils import get_column_letter

BG = "#FFFFFF"
WHITE = "#FFFFFF"
TEXT = "#37352F"
MUTED = "#787774"
BORDER = "#E9E9E7"
CELL_SHADOW = "#E9E9E7"
CELL_HOVER = "#F7F7F5"
CELL_HOVER_BORDER = "#D3D1CB"
PURPLE = "#2383E2"
PURPLE_DARK = "#0077D4"
PURPLE_PALE = "#E7F3F8"
GREEN = "#448361"
GREEN_PALE = "#EDF3EC"
RED = "#D44C47"
RED_PALE = "#FDEBEC"
AMBER = "#CB912F"
AMBER_PALE = "#FBF3DB"
CLASS_RE = re.compile(r"^\d{1,2}-[A-D]$", re.I)
TEACHER_RE = re.compile(r"^T\d{3,}$", re.I)
GROUP_MARK = "Group:"  # same marker as school_timetable_generator.GROUP_MARK


# Explain why a group period can't be edited like a single lesson.
def group_period_message(value):
    name = str(value or "").split("\n", 1)[0].strip() or "This group"
    return (name + " is a group period shared by its divisions and activity teachers, so it can't be changed "
            "here. Edit the group on the Requirements page and generate again.")


# Provide a GUI for editing timetable cells while keeping linked workbook sheets consistent.
class ExcelTimetableEditor:
    # Open the workbook, build the editor window, and connect UI actions to editing methods.
    def __init__(self, parent, workbook_path, on_saved=None, initial_sheet=None, locked_class=False, subject_first_mode=False):
        self.parent = parent
        self.path = os.path.abspath(workbook_path)
        self.on_saved = on_saved
        self.initial_sheet = initial_sheet
        self.locked_class = locked_class
        self.subject_first_mode = subject_first_mode
        self.dirty = False
        self.active_entry = None
        self.active_commit = None
        self.selected_cell = None
        self.cell_cards = {}
        self.teacher_sheets = {}
        self.tree = None
        try:
            self.workbook = load_workbook(self.path)
        except Exception as error:
            messagebox.showerror("Could not open Excel workbook", str(error), parent=parent)
            return
        self._refresh_teacher_sheet_map()

        self.window = tk.Toplevel(parent)
        self.window.title("Edit Excel timetable | SchoolFlow")
        sw, sh = self.window.winfo_screenwidth(), self.window.winfo_screenheight()
        self.window.geometry(f"{min(1600, max(1050, sw - 70))}x{min(1000, max(700, sh - 90))}")
        self.window.minsize(980, 640)
        self.window.configure(bg=BG)
        self.window.transient(parent)

        outer = tk.Frame(self.window, bg=BG)
        outer.pack(fill="both", expand=True, padx=18, pady=15)
        header = tk.Frame(outer, bg=BG)
        header.pack(fill="x", pady=(0, 12))
        tk.Label(header, text="Edit Excel timetable", bg=BG, fg=TEXT,
                 font=("Segoe UI", 19, "bold")).pack(anchor="w")
        if subject_first_mode and initial_sheet:
            header_help = (
                f"Focused on {initial_sheet}: double-click a period, choose a subject available in this grade, "
                "then choose one of its teachers. Teacher clashes are checked automatically."
            )
        else:
            header_help = (
                "Class sheets use a readable weekly grid. Double-click a slot to add or change a lesson; "
                "teacher clashes are checked automatically."
            )
        tk.Label(header, text=header_help,
                 bg=BG, fg=MUTED, font=("Segoe UI", 9), wraplength=1300,
                 justify="left").pack(anchor="w", pady=(4, 0))
        tk.Label(header, text=self.path, bg=BG, fg=PURPLE,
                 font=("Segoe UI", 8)).pack(anchor="w", pady=(5, 0))

        toolbar = tk.Frame(outer, bg=WHITE, highlightbackground=BORDER, highlightthickness=1)
        toolbar.pack(fill="x", pady=(0, 10))
        bar = tk.Frame(toolbar, bg=WHITE, padx=12, pady=10)
        bar.pack(fill="x")
        tk.Label(bar, text="WORKSHEET", bg=WHITE, fg=MUTED,
                 font=("Segoe UI", 8, "bold")).pack(side="left", padx=(0, 8))
        self.sheet_var = tk.StringVar(value=(initial_sheet if initial_sheet in self.workbook.sheetnames else self._default_sheet()))
        sheet_choices = [initial_sheet] if locked_class and initial_sheet in self.workbook.sheetnames else self.workbook.sheetnames
        self.sheet_box = ttk.Combobox(bar, textvariable=self.sheet_var,
                                      values=sheet_choices, state="readonly", width=25)
        self.sheet_box.pack(side="left", padx=(0, 14))
        self.sheet_box.bind("<<ComboboxSelected>>", self.change_sheet)
        style = ttk.Style(self.window)
        style.configure("Workbook.Treeview", background=WHITE, foreground=TEXT,
                        fieldbackground=WHITE, rowheight=30, borderwidth=0,
                        font=("Segoe UI", 9))
        style.configure("Workbook.Treeview.Heading", background="#F7F7F5",
                        foreground=TEXT, font=("Segoe UI", 9, "bold"), relief="flat", padding=8)
        style.map("Workbook.Treeview", background=[("selected", "#E3EEFA")])

        self.edit_button = tk.Button(bar, text="＋  Add / edit lesson",
            command=self.edit_selected_lesson, bg=PURPLE, fg=WHITE,
            activebackground=PURPLE_DARK, activeforeground=WHITE, relief="flat", bd=0,
            padx=14, pady=8, font=("Segoe UI", 9, "bold"), cursor="hand2")
        self.edit_button.pack(side="left", padx=(0, 7))
        self.clear_button = tk.Button(bar, text="Clear selected period",
            command=self.clear_selected_period, bg=RED_PALE, fg=RED,
            activebackground="#FAD7D5", activeforeground=RED, relief="flat", bd=0,
            padx=13, pady=8, font=("Segoe UI", 9, "bold"), cursor="hand2")
        self.clear_button.pack(side="left", padx=(0, 7))
        self.table_card = tk.Frame(outer, bg=WHITE, highlightbackground=BORDER, highlightthickness=1)
        self.table_card.pack(fill="both", expand=True)
        self.table_frame = tk.Frame(self.table_card, bg=WHITE)
        self.table_frame.pack(fill="both", expand=True, padx=1, pady=1)

        footer = tk.Frame(outer, bg=BG)
        footer.pack(fill="x", pady=(10, 0))
        self.status_var = tk.StringVar(value="Ready to edit.")
        tk.Label(footer, textvariable=self.status_var, bg=BG, fg=MUTED,
                 font=("Segoe UI", 9), anchor="w").pack(side="left", fill="x", expand=True)
        tk.Button(footer, text="Save workbook", command=self.save_changes,
            bg=PURPLE, fg=WHITE, activebackground=PURPLE_DARK, activeforeground=WHITE,
            relief="flat", bd=0, padx=19, pady=10, font=("Segoe UI", 10, "bold"),
            cursor="hand2").pack(side="right", padx=(8, 0))
        tk.Button(footer, text="Close", command=self.close_window,
            bg="#F1F1EF", fg=TEXT, activebackground="#E9E9E7", activeforeground=TEXT,
            relief="flat", bd=0, padx=18, pady=10, font=("Segoe UI", 9, "bold"),
            cursor="hand2").pack(side="right")
        tk.Label(outer, text="Class timetable edits update the linked teacher sheet and Requirements period count. Teacher availability is checked before saving.",
                 bg=BG, fg=MUTED, font=("Segoe UI", 8), wraplength=1350,
                 justify="left").pack(anchor="w", pady=(8, 0))
        self.window.protocol("WM_DELETE_WINDOW", self.close_window)
        self.load_sheet()
        self.window.grab_set()

    # Choose a sensible worksheet to show first, preferring a class timetable.
    def _default_sheet(self):
        for name in self.workbook.sheetnames:
            if CLASS_RE.match(name):
                return name
        for name in self.workbook.sheetnames:
            if name not in ("Summary", "Requirements", "Repeat Adjustments"):
                return name
        return self.workbook.sheetnames[0]

    # Return the worksheet selected in the worksheet dropdown.
    def _current_sheet(self):
        return self.workbook[self.sheet_var.get()]

    # Convert an Excel cell value into readable text for the interface.
    @staticmethod
    def _display_value(value):
        if value is None:
            return ""
        if isinstance(value, float) and value.is_integer():
            return str(int(value))
        return str(value)

    # Separate a timetable cell into its subject and teacher/class lines. A group
    # period (e.g. "AMD\nGroup: Music, Karate") is not a single lesson, so it reads
    # as ("", ""); use _is_group_cell to tell it apart from an empty period.
    @staticmethod
    def _split_lesson(value):
        if not value:
            return "", ""
        parts = str(value).split("\n", 1)
        second = parts[1].strip() if len(parts) > 1 else ""
        if second.startswith(GROUP_MARK):
            return "", ""
        return parts[0].strip(), second

    # Check whether a class cell holds a shared group period.
    @staticmethod
    def _is_group_cell(value):
        parts = str(value or "").split("\n", 1)
        return len(parts) > 1 and parts[1].strip().startswith(GROUP_MARK)

    # Check whether a worksheet name follows the grade-division format.
    @staticmethod
    def _is_class_sheet(name):
        return bool(CLASS_RE.match(name or ""))

    # Check whether a worksheet uses the generated teacher-sheet naming format.
    @staticmethod
    def _is_teacher_sheet(name):
        return bool(TEACHER_RE.match(name or ""))

    # Identify worksheets that contain class or teacher timetable grids.
    @classmethod
    def _is_timetable_sheet(cls, name):
        return cls._is_class_sheet(name) or cls._is_teacher_sheet(name)

    # Build a map from teacher names to their combined timetable worksheet names.
    def _refresh_teacher_sheet_map(self):
        self.teacher_sheets = {}
        if "Summary" in self.workbook.sheetnames:
            summary = self.workbook["Summary"]
            for row in range(2, summary.max_row + 1):
                teacher = self._display_value(summary.cell(row, 4).value).strip()
                sheet_name = self._display_value(summary.cell(row, 5).value).strip()
                if teacher and sheet_name in self.workbook.sheetnames:
                    self.teacher_sheets[teacher] = sheet_name
        for name in self.workbook.sheetnames:
            if self._is_teacher_sheet(name) and name not in self.teacher_sheets.values():
                title = self._display_value(self.workbook[name].cell(1, 1).value)
                if title.lower().startswith("teacher:"):
                    self.teacher_sheets[title.split(":", 1)[1].strip()] = name

    # Mark unsaved workbook changes and update the status message.
    def _set_dirty(self, message):
        self.dirty = True
        self.status_var.set(message)

    # Load the selected worksheet into the appropriate visual editor.
    def load_sheet(self):
        if self.active_commit is not None and not self.active_commit():
            return
        # Cancel pending per-cell animations before destroying the previous grid.
        for card in self.cell_cards.values():
            for timer_name in ("_hover_after", "_hover_check_after"):
                timer_id = getattr(card, timer_name, None)
                if timer_id is not None:
                    try:
                        card.after_cancel(timer_id)
                    except tk.TclError:
                        pass
        for widget in self.table_frame.winfo_children():
            widget.destroy()
        self.tree = None
        self.cell_cards = {}
        sheet = self._current_sheet()
        if self._is_timetable_sheet(sheet.title):
            self._load_timetable_grid(sheet)
        else:
            self._load_generic_sheet(sheet)
        is_class = self._is_class_sheet(sheet.title)
        self.edit_button.configure(state="normal" if is_class else "disabled")
        self.clear_button.configure(state="normal" if is_class else "disabled")
        suffix = "Unsaved changes" if self.dirty else "Ready"
        if self._is_teacher_sheet(sheet.title):
            info = "Teacher view · edit from a class timetable; this sheet updates automatically."
        elif is_class:
            info = "Double-click a lesson card or select one and press Add / edit lesson."
        elif sheet.title == "Requirements":
            info = "Edit input data here, or use class sheets for synchronized lesson changes."
        else:
            info = "Choose a class sheet to add or edit a lesson."
        self.status_var.set(f"{sheet.title} · {sheet.max_row} rows × {sheet.max_column} columns · {suffix} — {info}")

    # Handle a worksheet dropdown change and refresh the displayed sheet.
    def change_sheet(self, _event=None):
        self.selected_cell = None
        self.load_sheet()

    # Render a timetable worksheet as a grid of clickable day and period cards.
    def _load_timetable_grid(self, sheet):
        wrap = tk.Frame(self.table_frame, bg=WHITE)
        wrap.pack(fill="both", expand=True)
        canvas = tk.Canvas(wrap, bg=BG, highlightthickness=0)
        ybar = ttk.Scrollbar(wrap, orient="vertical", command=canvas.yview)
        xbar = ttk.Scrollbar(wrap, orient="horizontal", command=canvas.xview)
        canvas.configure(yscrollcommand=ybar.set, xscrollcommand=xbar.set)
        canvas.grid(row=0, column=0, sticky="nsew")
        ybar.grid(row=0, column=1, sticky="ns")
        xbar.grid(row=1, column=0, sticky="ew")
        wrap.rowconfigure(0, weight=1)
        wrap.columnconfigure(0, weight=1)
        inner = tk.Frame(canvas, bg=BG, padx=8, pady=8)
        window_id = canvas.create_window((0, 0), window=inner, anchor="nw")
        inner.bind("<Configure>", lambda _event: canvas.configure(scrollregion=canvas.bbox("all")))
        canvas.bind("<Configure>", lambda event: self._fit_canvas_width(canvas, inner, window_id, event))
        self.table_canvas = canvas
        self.table_inner = inner

        period_count = min(max(sheet.max_column - 1, 1), 12)
        banner = tk.Frame(inner, bg=WHITE, padx=8, pady=10)
        banner.grid(row=0, column=0, columnspan=period_count + 1, sticky="ew", pady=(0, 8))
        title = self._display_value(sheet.cell(1, 1).value) or sheet.title
        tk.Label(banner, text=title, bg=WHITE, fg=TEXT,
                 font=("Segoe UI", 16, "bold")).pack(side="left")
        subtext = ("Double-click any period to edit · teacher clashes checked" if self._is_class_sheet(sheet.title)
                   else "Read-only · changes are made from class timetable sheets")
        tk.Label(banner, text=subtext, bg=WHITE, fg=MUTED,
                 font=("Segoe UI", 9)).pack(side="right")

        tk.Label(inner, text="DAY", bg="#F1F1EF", fg=MUTED,
                 font=("Segoe UI", 8, "bold"), padx=10, pady=9,
                 width=13, anchor="w").grid(row=1, column=0, sticky="nsew", padx=3, pady=3)
        inner.grid_columnconfigure(0, minsize=118, weight=0)
        for p in range(period_count):
            label = self._display_value(sheet.cell(2, p + 2).value) or f"Period {p + 1}"
            tk.Label(inner, text=label.upper(), bg="#F1F1EF", fg=TEXT,
                     font=("Segoe UI", 8, "bold"), padx=8, pady=9).grid(
                row=1, column=p + 1, sticky="nsew", padx=3, pady=3)
            inner.grid_columnconfigure(p + 1, minsize=145, weight=1)

        day_rows = [r for r in range(3, sheet.max_row + 1) if self._display_value(sheet.cell(r, 1).value).strip()]
        for day_index, row_number in enumerate(day_rows):
            day = self._display_value(sheet.cell(row_number, 1).value)
            grid_row = day_index + 2
            tk.Label(inner, text=day, bg=WHITE, fg=TEXT,
                     font=("Segoe UI", 9, "bold"), padx=10, anchor="w").grid(
                row=grid_row, column=0, sticky="nsew", padx=3, pady=3)
            for p in range(period_count):
                col_number = p + 2
                value = self._display_value(sheet.cell(row_number, col_number).value)
                filled = bool(value)
                color = WHITE if filled else "#FFFFFF"
                # Keep the grid cell's geometry fixed. Its outer frame acts as a
                # soft shadow, while the inset face animates upward on hover.
                card = tk.Frame(inner, bg=CELL_SHADOW, highlightthickness=0,
                                height=82, width=145)
                card.grid(row=grid_row, column=p + 1, sticky="nsew", padx=3, pady=3)
                card.grid_propagate(False)
                face = tk.Frame(card, bg=color, highlightbackground=BORDER,
                                highlightthickness=1, padx=8, pady=7)
                face.place(x=0, y=2, relwidth=1, height=78)
                card._face = face
                card._base_color = color
                card._hover_progress = 0.0
                card._hovered = False
                card._hover_after = None
                card._hover_check_after = None
                line_one, line_two = self._split_lesson(value)
                if self._is_class_sheet(sheet.title):
                    if filled:
                        tk.Label(face, text=line_one, bg=color, fg=TEXT,
                                 font=("Segoe UI", 9, "bold"), anchor="w",
                                 justify="left", wraplength=125).pack(anchor="w", fill="x")
                        tk.Label(face, text=line_two or "Teacher not listed", bg=color, fg=MUTED,
                                 font=("Segoe UI", 8), anchor="w", justify="left",
                                 wraplength=125).pack(anchor="w", fill="x", pady=(5, 0))
                    else:
                        tk.Label(face, text="＋ Add lesson", bg=color, fg=PURPLE,
                                 font=("Segoe UI", 9, "bold")).pack(expand=True)
                        tk.Label(face, text="Free slot", bg=color, fg=MUTED,
                                 font=("Segoe UI", 8)).pack()
                    self._bind_card(card, row_number, col_number, sheet.title, True)
                else:
                    if filled:
                        tk.Label(face, text=line_one, bg=color, fg=TEXT,
                                 font=("Segoe UI", 9, "bold"), anchor="w",
                                 wraplength=125).pack(anchor="w", fill="x")
                        tk.Label(face, text=line_two, bg=color, fg=MUTED,
                                 font=("Segoe UI", 8), anchor="w",
                                 wraplength=125).pack(anchor="w", fill="x", pady=(5, 0))
                    else:
                        tk.Label(face, text="Free", bg=color, fg=MUTED,
                                 font=("Segoe UI", 9)).pack(expand=True)
                    self._bind_card(card, row_number, col_number, sheet.title, False)
                self.cell_cards[(row_number, col_number)] = card

    # Keep scrollable worksheet contents matched to the visible canvas width.
    @staticmethod
    def _fit_canvas_width(canvas, inner, window_id, event):
        canvas.itemconfigure(window_id, width=max(event.width, inner.winfo_reqwidth()))

    # Attach selection, editing, and animated-hover events to a timetable cell.
    def _bind_card(self, card, row_number, col_number, sheet_name, editable):
        def select(_event=None):
            self.selected_cell = (row_number, col_number)
            self._update_card_selection()
            sheet = self.workbook[sheet_name]
            day = self._display_value(sheet.cell(row_number, 1).value)
            period = self._display_value(sheet.cell(2, col_number).value)
            self.status_var.set(f"{sheet_name} · {day} · {period}" + (" · selected" if editable else " · view only"))

        def double_click(_event=None):
            select()
            if editable:
                self.edit_timetable_slot(row_number, col_number)

        def on_enter(_event=None):
            # Moving from a card into one of its text labels should not cancel hover.
            timer_id = getattr(card, "_hover_check_after", None)
            if timer_id is not None:
                try:
                    card.after_cancel(timer_id)
                except tk.TclError:
                    pass
                card._hover_check_after = None
            card._hovered = True
            self._animate_card_hover(card, 1.0)

        def on_leave(_event=None):
            # Check after the Enter event for the next child has had time to run.
            timer_id = getattr(card, "_hover_check_after", None)
            if timer_id is not None:
                try:
                    card.after_cancel(timer_id)
                except tk.TclError:
                    pass
            def check_pointer():
                card._hover_check_after = None
                try:
                    pointer_x, pointer_y = card.winfo_pointerx(), card.winfo_pointery()
                    left, top = card.winfo_rootx(), card.winfo_rooty()
                    inside = (left <= pointer_x < left + card.winfo_width() and
                              top <= pointer_y < top + card.winfo_height())
                except tk.TclError:
                    return
                if not inside:
                    card._hovered = False
                    self._animate_card_hover(card, 0.0)
            try:
                card._hover_check_after = card.after(35, check_pointer)
            except tk.TclError:
                pass

        # Bind the full card hierarchy so mouse events work over lesson text too.
        widgets = []
        def collect_widgets(parent):
            widgets.append(parent)
            for child in parent.winfo_children():
                collect_widgets(child)
        collect_widgets(card)
        for widget in widgets:
            widget.bind("<Button-1>", select, add="+")
            widget.bind("<Double-Button-1>", double_click, add="+")
            widget.bind("<Enter>", on_enter, add="+")
            widget.bind("<Leave>", on_leave, add="+")

    # Return the cell's base colour, independent of its current hover/selection state.
    @staticmethod
    def _normal_card_color(card):
        return getattr(card, "_base_color", "#FFFFFF")

    # Blend two hex colours so the card can transition smoothly while hovering.
    @staticmethod
    def _blend_card_color(first, second, amount):
        amount = max(0.0, min(1.0, amount))
        channels = []
        for offset in (1, 3, 5):
            a = int(first[offset:offset + 2], 16)
            b = int(second[offset:offset + 2], 16)
            channels.append(round(a + (b - a) * amount))
        return "#" + "".join(f"{channel:02X}" for channel in channels)

    # Animate the face's vertical offset and tint without changing the grid's size.
    def _animate_card_hover(self, card, target):
        timer_id = getattr(card, "_hover_after", None)
        if timer_id is not None:
            try:
                card.after_cancel(timer_id)
            except tk.TclError:
                pass
            card._hover_after = None
        start = float(getattr(card, "_hover_progress", 0.0))
        target = float(target)
        if abs(start - target) < 0.001:
            card._hover_progress = target
            self._render_card_appearance(card)
            return
        step = {"value": 0}
        total_steps = 6

        def animate_step():
            try:
                if not card.winfo_exists():
                    return
                step["value"] += 1
                linear = min(1.0, step["value"] / total_steps)
                eased = linear * linear * (3 - 2 * linear)
                card._hover_progress = start + (target - start) * eased
                self._render_card_appearance(card)
                if step["value"] < total_steps:
                    card._hover_after = card.after(14, animate_step)
                else:
                    card._hover_after = None
            except tk.TclError:
                return

        card._hover_after = card.after(0, animate_step)

    # Apply selection and animated-hover styles while keeping border thickness fixed.
    def _render_card_appearance(self, card):
        progress = float(getattr(card, "_hover_progress", 0.0))
        selected = self.selected_cell is not None and self.cell_cards.get(self.selected_cell) is card
        base = PURPLE_PALE if selected else self._normal_card_color(card)
        hover = "#E7F3F8" if selected else CELL_HOVER
        color = self._blend_card_color(base, hover, progress)
        border_base = PURPLE if selected else BORDER
        border_hover = PURPLE if selected else CELL_HOVER_BORDER
        border = self._blend_card_color(border_base, border_hover, progress)
        try:
            face = card._face
            face.configure(bg=color, highlightbackground=border, highlightthickness=1)
            # A 2 px movement inside the fixed outer card produces a lift/shadow effect
            # without changing neighbouring cell positions or causing hover jitter.
            face.place_configure(y=2 - round(2 * progress))
            for child in face.winfo_children():
                child.configure(bg=color)
        except (AttributeError, tk.TclError):
            pass

    # Refresh each cell's appearance after the selected timetable slot changes.
    def _update_card_selection(self):
        for card in self.cell_cards.values():
            self._render_card_appearance(card)

    # Display a non-timetable worksheet in a generic row-and-column table.
    def _load_generic_sheet(self, sheet):
        max_row = max(sheet.max_row, 1)
        max_col = min(max(sheet.max_column, 1), 80)
        columns = ["col_" + str(i) for i in range(1, max_col + 1)]
        container = tk.Frame(self.table_frame, bg=WHITE)
        container.pack(fill="both", expand=True)
        self.tree = ttk.Treeview(container, columns=columns, show="tree headings",
                                 style="Workbook.Treeview", selectmode="browse")
        self.tree.heading("#0", text="Row")
        self.tree.column("#0", width=62, minwidth=52, stretch=False, anchor="center")
        ybar = ttk.Scrollbar(container, orient="vertical", command=self.tree.yview)
        xbar = ttk.Scrollbar(container, orient="horizontal", command=self.tree.xview)
        self.tree.configure(yscrollcommand=ybar.set, xscrollcommand=xbar.set)
        self.tree.grid(row=0, column=0, sticky="nsew")
        ybar.grid(row=0, column=1, sticky="ns")
        xbar.grid(row=1, column=0, sticky="ew")
        container.rowconfigure(0, weight=1)
        container.columnconfigure(0, weight=1)
        for i, column_id in enumerate(columns, start=1):
            heading = self._display_value(sheet.cell(1, i).value) or get_column_letter(i)
            self.tree.heading(column_id, text=heading)
            self.tree.column(column_id, width=155 if i > 1 else 110, minwidth=85,
                             stretch=True, anchor="center" if i > 1 else "w")
        for row in range(1, max_row + 1):
            values = [self._display_value(sheet.cell(row, col).value) for col in range(1, max_col + 1)]
            self.tree.insert("", "end", iid=str(row), text=str(row), values=values,
                             tags=("even" if row % 2 else "odd",))
        self.tree.tag_configure("even", background=WHITE)
        self.tree.tag_configure("odd", background="#FBFBFA")
        self.tree.bind("<Double-1>", self.begin_cell_edit)
        self.tree.bind("<ButtonRelease-1>", self.record_selected_cell, add="+")

    # Find Requirements rows that belong to the selected class division.
    def _requirement_rows_for_class(self, class_name):
        if "Requirements" not in self.workbook.sheetnames:
            return []
        grade, division = class_name.split("-", 1)
        sheet = self.workbook["Requirements"]
        result = []
        for row in range(2, sheet.max_row + 1):
            same_class = (self._display_value(sheet.cell(row, 1).value).strip() == grade and
                          self._display_value(sheet.cell(row, 2).value).strip().upper() == division.upper())
            subject = self._display_value(sheet.cell(row, 3).value).strip()
            teacher = self._display_value(sheet.cell(row, 4).value).strip()
            if same_class and subject and teacher:
                result.append({"row": row, "subject": subject, "teacher": teacher,
                               "periods": int(sheet.cell(row, 5).value or 0)})
        return result

    # Extract the grade number from a class name such as 12-A.
    def _grade_for_class(self, class_name):
        return class_name.split("-", 1)[0].strip()

    # List the subjects assigned anywhere within the selected grade.
    def _subjects_for_grade(self, class_name):
        """List subjects already assigned in this grade across all divisions."""
        grade = self._grade_for_class(class_name)
        subjects = {}
        if "Requirements" in self.workbook.sheetnames:
            sheet = self.workbook["Requirements"]
            for row in range(2, sheet.max_row + 1):
                row_grade = self._display_value(sheet.cell(row, 1).value).strip()
                subject = self._display_value(sheet.cell(row, 3).value).strip()
                if row_grade == grade and subject:
                    subjects.setdefault(subject.casefold(), subject)
        # Also include subjects already present in the timetable in case the
        # Requirements sheet is missing a newly added lesson.
        for name in self.workbook.sheetnames:
            if not self._is_class_sheet(name) or self._grade_for_class(name) != grade:
                continue
            sheet = self.workbook[name]
            for row in range(3, sheet.max_row + 1):
                for col in range(2, sheet.max_column + 1):
                    subject, _teacher = self._split_lesson(self._display_value(sheet.cell(row, col).value))
                    if subject:
                        subjects.setdefault(subject.casefold(), subject)
        return sorted(subjects.values(), key=str.casefold)

    # List teachers who are assigned to subjects in the selected grade.
    def _teachers_for_grade(self, class_name):
        """Return teachers assigned to any subject in the selected grade."""
        grade = self._grade_for_class(class_name)
        teachers = {}
        if "Requirements" in self.workbook.sheetnames:
            sheet = self.workbook["Requirements"]
            for row in range(2, sheet.max_row + 1):
                row_grade = self._display_value(sheet.cell(row, 1).value).strip()
                teacher = self._display_value(sheet.cell(row, 4).value).strip()
                if row_grade == grade and teacher:
                    teachers.setdefault(teacher.casefold(), teacher)
        for name in self.workbook.sheetnames:
            if not self._is_class_sheet(name) or self._grade_for_class(name) != grade:
                continue
            sheet = self.workbook[name]
            for row in range(3, sheet.max_row + 1):
                for col in range(2, sheet.max_column + 1):
                    _subject, teacher = self._split_lesson(self._display_value(sheet.cell(row, col).value))
                    if teacher:
                        teachers.setdefault(teacher.casefold(), teacher)
        return sorted(teachers.values(), key=str.casefold)

    # Find teachers assigned to one subject across all divisions in the grade.
    def _teachers_for_subject_in_grade(self, class_name, subject):
        """Return every teacher assigned to this subject in any division of the grade."""
        subject = subject.strip().casefold()
        if not subject:
            return []
        grade = self._grade_for_class(class_name)
        teachers = {}
        if "Requirements" in self.workbook.sheetnames:
            sheet = self.workbook["Requirements"]
            for row in range(2, sheet.max_row + 1):
                row_grade = self._display_value(sheet.cell(row, 1).value).strip()
                row_subject = self._display_value(sheet.cell(row, 3).value).strip().casefold()
                teacher = self._display_value(sheet.cell(row, 4).value).strip()
                if row_grade == grade and row_subject == subject and teacher:
                    teachers.setdefault(teacher.casefold(), teacher)
        # Timetable cells are also checked, so manually added lessons are not
        # ignored even if the Requirements sheet has not been refreshed yet.
        for name in self.workbook.sheetnames:
            if not self._is_class_sheet(name) or self._grade_for_class(name) != grade:
                continue
            sheet = self.workbook[name]
            for row in range(3, sheet.max_row + 1):
                for col in range(2, sheet.max_column + 1):
                    row_subject, teacher = self._split_lesson(self._display_value(sheet.cell(row, col).value))
                    if row_subject.casefold() == subject and teacher:
                        teachers.setdefault(teacher.casefold(), teacher)
        return sorted(teachers.values(), key=str.casefold)

    # Return unique teacher-and-subject pairs found in the grade.
    def _teacher_subject_pairs_for_grade(self, class_name):
        """Return unique teacher/subject assignments used in the selected grade."""
        grade = self._grade_for_class(class_name)
        pairs = {}

        # Requirements is the main source of subject-teacher assignments.
        if "Requirements" in self.workbook.sheetnames:
            sheet = self.workbook["Requirements"]
            for row in range(2, sheet.max_row + 1):
                row_grade = self._display_value(sheet.cell(row, 1).value).strip()
                subject = self._display_value(sheet.cell(row, 3).value).strip()
                teacher = self._display_value(sheet.cell(row, 4).value).strip()
                if row_grade == grade and subject and teacher:
                    key = (teacher.casefold(), subject.casefold())
                    pairs.setdefault(key, {"teacher": teacher, "subject": subject})

        # Also scan class timetables so manual/recent assignments are included.
        for name in self.workbook.sheetnames:
            if not self._is_class_sheet(name) or self._grade_for_class(name) != grade:
                continue
            sheet = self.workbook[name]
            for row in range(3, sheet.max_row + 1):
                for col in range(2, sheet.max_column + 1):
                    subject, teacher = self._split_lesson(
                        self._display_value(sheet.cell(row, col).value))
                    if subject and teacher:
                        key = (teacher.casefold(), subject.casefold())
                        pairs.setdefault(key, {"teacher": teacher, "subject": subject})

        return sorted(pairs.values(),
                      key=lambda item: (item["teacher"].casefold(), item["subject"].casefold()))

    # Collect the teacher names known to the workbook.
    def _all_teachers(self):
        return sorted(self.teacher_sheets.keys(), key=str.casefold)

    # Find every subject currently assigned to a teacher across the school.
    def _subjects_assigned_to_teacher(self, teacher):
        """Return the subjects currently assigned to this teacher across the whole school."""
        teacher_key = teacher.strip().casefold()
        if not teacher_key:
            return set()
        subjects = set()
        if "Requirements" in self.workbook.sheetnames:
            sheet = self.workbook["Requirements"]
            for row in range(2, sheet.max_row + 1):
                row_teacher = self._display_value(sheet.cell(row, 4).value).strip()
                row_subject = self._display_value(sheet.cell(row, 3).value).strip()
                if row_teacher.casefold() == teacher_key and row_subject:
                    subjects.add(row_subject)
        # Class sheets are also checked in case a timetable was edited manually.
        for name in self.workbook.sheetnames:
            if not self._is_class_sheet(name):
                continue
            sheet = self.workbook[name]
            for row in range(3, sheet.max_row + 1):
                for col in range(2, sheet.max_column + 1):
                    subject, row_teacher = self._split_lesson(self._display_value(sheet.cell(row, col).value))
                    if row_teacher.casefold() == teacher_key and subject:
                        subjects.add(subject)
        return subjects

    # Check whether assigning a subject would violate the one-subject-per-teacher rule.
    def _teacher_subject_conflict(self, teacher, subject):
        """Return existing subject if teacher is already assigned a different subject."""
        if not teacher or not subject:
            return None
        subjects = self._subjects_assigned_to_teacher(teacher)
        other_subjects = [item for item in subjects if item.casefold() != subject.strip().casefold()]
        return sorted(other_subjects, key=str.casefold)[0] if other_subjects else None

    # Find out whether a teacher already has a lesson in the requested time slot.
    def _teacher_clash(self, teacher, class_name, row_number, col_number):
        """Return (other class, subject) if this teacher already has a lesson."""
        if not teacher:
            return None
        sheet_name = self.teacher_sheets.get(teacher)
        if sheet_name in self.workbook.sheetnames:
            value = self._display_value(self.workbook[sheet_name].cell(row_number, col_number).value)
            if value:
                other_class, other_subject = self._split_lesson(value)
                if other_class and other_class != class_name:
                    return other_class, other_subject
        # Check class sheets too, so stale or manually edited teacher sheets don't hide a clash.
        for name in self.workbook.sheetnames:
            if not self._is_class_sheet(name) or name == class_name:
                continue
            other = self._display_value(self.workbook[name].cell(row_number, col_number).value)
            other_subject, assigned_teacher = self._split_lesson(other)
            if assigned_teacher and assigned_teacher.casefold() == teacher.casefold():
                return name, other_subject
        return None

    # Determine how many period columns the current timetable contains.
    def _period_count(self):
        for name in self.workbook.sheetnames:
            if self._is_class_sheet(name):
                return max(1, self.workbook[name].max_column - 1)
        return 7

    # Read the day labels used by the current timetable grid.
    def _day_names(self):
        for name in self.workbook.sheetnames:
            if self._is_class_sheet(name):
                sheet = self.workbook[name]
                return [self._display_value(sheet.cell(row, 1).value)
                        for row in range(3, sheet.max_row + 1) if sheet.cell(row, 1).value]
        return ["Monday", "Tuesday", "Wednesday", "Thursday", "Friday"]

    # Create or populate the combined timetable sheet for a teacher when required.
    def _populate_teacher_sheet_if_missing(self, teacher):
        if teacher in self.teacher_sheets and self.teacher_sheets[teacher] in self.workbook.sheetnames:
            return self.teacher_sheets[teacher]
        summary = self.workbook["Summary"] if "Summary" in self.workbook.sheetnames else self.workbook.create_sheet("Summary", 0)
        if not summary.cell(1, 4).value:
            summary.cell(1, 4, "Teacher")
            summary.cell(1, 5, "Sheet")
        numbers = [int(name[1:]) for name in self.workbook.sheetnames if self._is_teacher_sheet(name)]
        sheet_name = "T" + str(max(numbers + [0]) + 1).zfill(3)
        sheet = self.workbook.create_sheet(sheet_name)
        sheet.cell(1, 1, "Teacher: " + teacher)
        sheet.merge_cells(start_row=1, start_column=1, end_row=1, end_column=self._period_count() + 1)
        sheet.cell(2, 1, "Day")
        for p in range(self._period_count()):
            sheet.cell(2, p + 2, "Period " + str(p + 1))
        for i, day in enumerate(self._day_names()):
            sheet.cell(i + 3, 1, day)
        next_row = max(2, summary.max_row) + 1
        summary.cell(next_row, 4, teacher)
        summary.cell(next_row, 5, sheet_name)
        self.teacher_sheets[teacher] = sheet_name
        return sheet_name

    # Locate the Requirements row for a specific class, subject, and teacher assignment.
    def _find_requirement_row(self, class_name, subject, teacher):
        if "Requirements" not in self.workbook.sheetnames:
            return None
        grade, division = class_name.split("-", 1)
        sheet = self.workbook["Requirements"]
        for row in range(2, sheet.max_row + 1):
            correct_class = (self._display_value(sheet.cell(row, 1).value).strip() == grade and
                             self._display_value(sheet.cell(row, 2).value).strip().upper() == division.upper())
            correct_subject = self._display_value(sheet.cell(row, 3).value).strip().casefold() == subject.strip().casefold()
            correct_teacher = self._display_value(sheet.cell(row, 4).value).strip().casefold() == teacher.strip().casefold()
            if correct_class and correct_subject and correct_teacher:
                return row
        return None

    # Adjust the weekly period count for an assignment in the Requirements sheet.
    def _change_requirement_count(self, class_name, subject, teacher, delta):
        if "Requirements" not in self.workbook.sheetnames:
            return
        sheet = self.workbook["Requirements"]
        row = self._find_requirement_row(class_name, subject, teacher)
        if row is not None:
            count = int(sheet.cell(row, 5).value or 0) + delta
            if count <= 0:
                sheet.delete_rows(row, 1)
            else:
                sheet.cell(row, 5, count)
        elif delta > 0:
            grade, division = class_name.split("-", 1)
            sheet.append([int(grade), division, subject, teacher, delta])

    # Remove a lesson from the matching teacher timetable slot.
    def _clear_teacher_slot(self, teacher, class_name, subject, row_number, col_number):
        sheet_name = self.teacher_sheets.get(teacher)
        if not sheet_name or sheet_name not in self.workbook.sheetnames:
            return
        cell = self.workbook[sheet_name].cell(row_number, col_number)
        assigned_class, assigned_subject = self._split_lesson(self._display_value(cell.value))
        if assigned_class == class_name and assigned_subject.casefold() == subject.casefold():
            cell.value = None

    # Apply a cell edit and update class, teacher, and requirement data together.
    def _apply_lesson_change(self, class_name, row_number, col_number, subject, teacher):
        class_sheet = self.workbook[class_name]
        cell = class_sheet.cell(row_number, col_number)
        if self._is_group_cell(cell.value):
            raise ValueError(group_period_message(self._display_value(cell.value)))
        old_subject, old_teacher = self._split_lesson(self._display_value(cell.value))
        subject, teacher = subject.strip(), teacher.strip()
        if subject and teacher:
            subject_conflict = self._teacher_subject_conflict(teacher, subject)
            if subject_conflict:
                raise ValueError(
                    f"{teacher} is assigned to {subject_conflict} and can only teach one subject. "
                    f"Choose a teacher assigned to {subject}."
                )
            clash = self._teacher_clash(teacher, class_name, row_number, col_number)
            if clash:
                other_class, other_subject = clash
                day = self._display_value(class_sheet.cell(row_number, 1).value)
                period = self._display_value(class_sheet.cell(2, col_number).value)
                raise ValueError(f"CLASH DETECTED: {teacher} is already teaching {other_subject or 'a lesson'} in Class {other_class} during {day}, {period}. Choose another teacher or period.")

        same_assignment = (old_subject.casefold(), old_teacher.casefold()) == (subject.casefold(), teacher.casefold())
        if old_subject and old_teacher and not same_assignment:
            self._clear_teacher_slot(old_teacher, class_name, old_subject, row_number, col_number)
            self._change_requirement_count(class_name, old_subject, old_teacher, -1)
        elif (not subject or not teacher) and old_subject and old_teacher:
            self._clear_teacher_slot(old_teacher, class_name, old_subject, row_number, col_number)
            self._change_requirement_count(class_name, old_subject, old_teacher, -1)

        if subject and teacher:
            if not same_assignment:
                self._change_requirement_count(class_name, subject, teacher, 1)
            cell.value = subject + "\n" + teacher
            teacher_sheet_name = self._populate_teacher_sheet_if_missing(teacher)
            self.workbook[teacher_sheet_name].cell(row_number, col_number, class_name + "\n" + subject)
            message = f"Assigned {subject} to {class_name}; no teacher clash found."
        else:
            cell.value = None
            message = "Period cleared; teacher timetable and Requirements count updated."

        self.dirty = True
        self._refresh_teacher_sheet_map()
        self.selected_cell = (row_number, col_number)
        self.load_sheet()
        self.status_var.set(message + " Save the workbook to keep changes.")

    # Edit a class slot by choosing an available grade subject before choosing its teacher.
    def _edit_timetable_slot_subject_first(self, row_number, col_number):
        """Edit a single-class timetable by selecting a grade subject first."""
        sheet = self._current_sheet()
        class_name = sheet.title
        grade = self._grade_for_class(class_name)
        day = self._display_value(sheet.cell(row_number, 1).value)
        period = self._display_value(sheet.cell(2, col_number).value)
        old_subject, old_teacher = self._split_lesson(
            self._display_value(sheet.cell(row_number, col_number).value))

        subjects = self._subjects_for_grade(class_name)
        if old_subject and old_subject not in subjects:
            subjects.append(old_subject)
        subjects = sorted(set(subjects), key=str.casefold)
        if not subjects:
            messagebox.showwarning(
                "No subjects available",
                f"No subjects are configured for Grade {grade}. Add subject assignments to the Requirements sheet first.",
                parent=self.window,
            )
            return

        dialog = tk.Toplevel(self.window)
        dialog.title("Edit lesson · " + class_name)
        dialog.geometry("650x440")
        dialog.minsize(580, 410)
        dialog.configure(bg=BG)
        dialog.transient(self.window)
        dialog.grab_set()

        body = tk.Frame(dialog, bg=BG, padx=24, pady=20)
        body.pack(fill="both", expand=True)
        tk.Label(body, text="Edit class lesson", bg=BG, fg=TEXT,
                 font=("Segoe UI", 17, "bold")).pack(anchor="w")
        tk.Label(body, text=f"{class_name}  ·  {day}  ·  {period}", bg=BG,
                 fg=MUTED, font=("Segoe UI", 10)).pack(anchor="w", pady=(4, 17))

        tk.Label(body, text=f"SUBJECT · AVAILABLE IN GRADE {grade}", bg=BG,
                 fg=MUTED, font=("Segoe UI", 8, "bold")).pack(anchor="w")
        subject_var = tk.StringVar(value=old_subject if old_subject in subjects else subjects[0])
        subject_box = ttk.Combobox(body, textvariable=subject_var, values=subjects,
                                   state="readonly", font=("Segoe UI", 10))
        subject_box.pack(fill="x", ipady=5, pady=(5, 13))

        tk.Label(body, text="TEACHER · TEACHES THE SELECTED SUBJECT", bg=BG,
                 fg=MUTED, font=("Segoe UI", 8, "bold")).pack(anchor="w")
        teacher_var = tk.StringVar(value="")
        teacher_box = ttk.Combobox(body, textvariable=teacher_var, values=(),
                                   state="readonly", font=("Segoe UI", 10))
        teacher_box.pack(fill="x", ipady=5, pady=(5, 8))

        hint_var = tk.StringVar(value="Choose a subject to see matching teachers.")
        tk.Label(body, textvariable=hint_var, bg=BG, fg=MUTED,
                 font=("Segoe UI", 8), anchor="w", justify="left",
                 wraplength=590).pack(fill="x", pady=(0, 7))

        status_card = tk.Frame(body, bg=WHITE, highlightbackground=BORDER,
                               highlightthickness=1, padx=12, pady=10)
        status_card.pack(fill="x", pady=(0, 13))
        status_label = tk.Label(status_card, text="Select a subject and teacher to check availability.",
                                bg=WHITE, fg=MUTED, font=("Segoe UI", 9),
                                anchor="w", justify="left", wraplength=570)
        status_label.pack(fill="x")
        teacher_by_label = {}

        def update_teachers(_event=None):
            subject = subject_var.get().strip()
            teachers = self._teachers_for_subject_in_grade(class_name, subject)
            if old_teacher and old_subject.casefold() == subject.casefold() and old_teacher not in teachers:
                teachers.append(old_teacher)
            entries = []
            teacher_by_label.clear()
            available = 0
            busy = 0
            for teacher in sorted(set(teachers), key=str.casefold):
                clash = self._teacher_clash(teacher, class_name, row_number, col_number)
                subject_conflict = self._teacher_subject_conflict(teacher, subject)
                if clash:
                    label = f"{teacher}  ·  Busy in {clash[0]}"
                    busy += 1
                elif subject_conflict:
                    label = f"{teacher}  ·  Assigned to {subject_conflict}"
                    busy += 1
                else:
                    label = f"{teacher}  ·  Available"
                    available += 1
                entries.append(label)
                teacher_by_label[label] = teacher

            entries.sort(key=lambda label: (
                0 if label.endswith("· Available") else 1, label.casefold()
            ))
            teacher_box.configure(values=entries)
            preferred = next((label for label in entries
                              if teacher_by_label[label].casefold() == old_teacher.casefold()
                              and subject.casefold() == old_subject.casefold()), None)
            chosen = preferred or next((label for label in entries if label.endswith("· Available")), None)
            if chosen is None and entries:
                chosen = entries[0]
            teacher_var.set(chosen or "")
            hint_var.set(
                f"Grade {grade} · {len(entries)} matching teachers · {available} available · {busy} busy"
                if entries else f"No teachers are assigned to {subject} in Grade {grade}. Add an assignment in Requirements first."
            )
            check_status()

        def check_status(_event=None):
            subject = subject_var.get().strip()
            teacher = teacher_by_label.get(teacher_var.get().strip(), "")
            if not subject or not teacher:
                status_label.configure(
                    text=f"No teacher is currently assigned to this subject in Grade {grade}. Choose another subject or update Requirements.",
                    bg=AMBER_PALE, fg=AMBER,
                )
                status_card.configure(bg=AMBER_PALE, highlightbackground=AMBER)
                return
            clash = self._teacher_clash(teacher, class_name, row_number, col_number)
            subject_conflict = self._teacher_subject_conflict(teacher, subject)
            if clash:
                status_label.configure(
                    text=f"CLASH DETECTED · {teacher} is teaching {clash[1] or 'a lesson'} in Class {clash[0]} at this time. Choose another teacher.",
                    bg=RED_PALE, fg=RED,
                )
                status_card.configure(bg=RED_PALE, highlightbackground=RED)
            elif subject_conflict:
                status_label.configure(
                    text=f"SUBJECT RULE · {teacher} is assigned to {subject_conflict}, so they cannot teach {subject}.",
                    bg=RED_PALE, fg=RED,
                )
                status_card.configure(bg=RED_PALE, highlightbackground=RED)
            else:
                status_label.configure(
                    text=f"AVAILABLE · {teacher} can teach {subject} during this period.",
                    bg=GREEN_PALE, fg=GREEN,
                )
                status_card.configure(bg=GREEN_PALE, highlightbackground=GREEN)

        def save_lesson():
            subject = subject_var.get().strip()
            teacher = teacher_by_label.get(teacher_var.get().strip(), "")
            if not subject or not teacher:
                messagebox.showwarning(
                    "Choose a teacher",
                    "Select a teacher who is assigned to the selected subject in this grade.",
                    parent=dialog,
                )
                return
            clash = self._teacher_clash(teacher, class_name, row_number, col_number)
            if clash:
                messagebox.showerror(
                    "Teacher clash",
                    f"{teacher} is already teaching {clash[1] or 'a lesson'} in Class {clash[0]} during {day}, {period}.",
                    parent=dialog,
                )
                return
            try:
                self._apply_lesson_change(class_name, row_number, col_number, subject, teacher)
            except ValueError as error:
                messagebox.showerror("Could not save lesson", str(error), parent=dialog)
                return
            dialog.destroy()

        subject_box.bind("<<ComboboxSelected>>", update_teachers)
        teacher_box.bind("<<ComboboxSelected>>", check_status)
        update_teachers()

        actions = tk.Frame(body, bg=BG)
        actions.pack(side="bottom", fill="x", pady=(4, 0))
        tk.Button(actions, text="Cancel", command=dialog.destroy,
                  bg="#F1F1EF", fg=TEXT, relief="flat", padx=16, pady=8,
                  cursor="hand2").pack(side="right")
        tk.Button(actions, text="Save lesson", command=save_lesson,
                  bg=PURPLE, fg=WHITE, activebackground=PURPLE_DARK,
                  relief="flat", padx=18, pady=8, font=("Segoe UI", 9, "bold"),
                  cursor="hand2").pack(side="right", padx=(0, 8))
        dialog.update_idletasks()
        dialog.grab_set()

    # Open the editing workflow for the currently selected timetable cell.
    def edit_selected_lesson(self):
        sheet = self._current_sheet()
        if not self._is_class_sheet(sheet.title):
            messagebox.showinfo("Choose a class sheet", "Select a class timetable such as 1-A to edit lessons.", parent=self.window)
            return
        if not self.selected_cell:
            messagebox.showinfo("Select a period", "Click a period card first, then choose Add / edit lesson.", parent=self.window)
            return
        row, col = self.selected_cell
        if row < 3 or col < 2:
            messagebox.showinfo("Select a period", "Select a period card in the timetable grid.", parent=self.window)
            return
        self.edit_timetable_slot(row, col)

    # Show the dialog used to choose or change a lesson in one time slot.
    def edit_timetable_slot(self, row_number, col_number):
        sheet = self._current_sheet()
        if not self._is_class_sheet(sheet.title):
            messagebox.showinfo("Read-only teacher timetable", "Switch to a class timetable to change a lesson.", parent=self.window)
            return
        if row_number < 3 or col_number < 2 or row_number > sheet.max_row or col_number > sheet.max_column:
            return

        if self.subject_first_mode:
            return self._edit_timetable_slot_subject_first(row_number, col_number)

        class_name = sheet.title
        grade = self._grade_for_class(class_name)
        day = self._display_value(sheet.cell(row_number, 1).value)
        period = self._display_value(sheet.cell(2, col_number).value)
        old_subject, old_teacher = self._split_lesson(
            self._display_value(sheet.cell(row_number, col_number).value))

        # Use existing teacher/subject pairs. The subject is display-only; the
        # selected pair determines both the teacher and the subject automatically.
        pairs = self._teacher_subject_pairs_for_grade(class_name)
        old_pair_key = (old_teacher.casefold(), old_subject.casefold())
        if old_teacher and old_subject and not any(
                (p["teacher"].casefold(), p["subject"].casefold()) == old_pair_key for p in pairs):
            pairs.append({"teacher": old_teacher, "subject": old_subject})
            pairs.sort(key=lambda item: (item["teacher"].casefold(), item["subject"].casefold()))

        dialog = tk.Toplevel(self.window)
        dialog.title("Edit lesson · " + class_name)
        dialog.geometry("660x470")
        dialog.minsize(600, 450)
        dialog.configure(bg=BG)
        dialog.transient(self.window)
        dialog.grab_set()
        body = tk.Frame(dialog, bg=BG, padx=24, pady=20)
        body.pack(fill="both", expand=True)
        tk.Label(body, text="Edit lesson", bg=BG, fg=TEXT,
                 font=("Segoe UI", 17, "bold")).pack(anchor="w")
        tk.Label(body, text=f"{class_name}  ·  {day}  ·  {period}", bg=BG,
                 fg=MUTED, font=("Segoe UI", 10)).pack(anchor="w", pady=(4, 16))

        pair_labels = []
        pair_by_label = {}
        # Show availability directly in the dropdown so users can avoid busy teachers.
        for item in pairs:
            teacher = item["teacher"]
            subject = item["subject"]
            clash = self._teacher_clash(teacher, class_name, row_number, col_number)
            if clash:
                availability = f"Busy in {clash[0]}"
            else:
                availability = "Available"
            label = f"{teacher}  —  {subject}   ·   {availability}"
            # Keep labels unique even if the input data contains duplicate display names.
            if label in pair_by_label:
                continue
            pair_labels.append(label)
            pair_by_label[label] = (teacher, subject)

        # Put usable options first, with busy assignments still visible for inspection.
        pair_labels.sort(key=lambda label: (0 if "·   Available" in label else 1, label.casefold()))

        tk.Label(body, text="TEACHER · CHOOSE AN ASSIGNMENT", bg=BG,
                 fg=MUTED, font=("Segoe UI", 8, "bold")).pack(anchor="w")
        assignment_var = tk.StringVar(value="")
        assignment_box = ttk.Combobox(body, textvariable=assignment_var, values=pair_labels,
                                      state="readonly", font=("Segoe UI", 10))
        assignment_box.pack(fill="x", ipady=5, pady=(5, 6))

        tk.Label(body, text="SUBJECT · AUTO-FILLED FROM THE SELECTED TEACHER", bg=BG,
                 fg=MUTED, font=("Segoe UI", 8, "bold")).pack(anchor="w", pady=(7, 0))
        subject_var = tk.StringVar(value=old_subject)
        subject_box = ttk.Entry(body, textvariable=subject_var, state="readonly",
                                font=("Segoe UI", 10))
        subject_box.pack(fill="x", ipady=5, pady=(5, 9))

        availability_row = tk.Frame(body, bg=BG)
        availability_row.pack(fill="x", pady=(0, 10))
        teacher_hint = tk.Label(
            availability_row,
            text="Checking teacher availability…",
            bg=BG, fg=MUTED, font=("Segoe UI", 8), justify="left", anchor="w")
        teacher_hint.pack(side="left", fill="x", expand=True)

        clash_card = tk.Frame(body, bg=WHITE, highlightbackground=BORDER,
                              highlightthickness=1, padx=12, pady=10)
        clash_card.pack(fill="x", pady=(0, 14))
        clash_label = tk.Label(clash_card, text="Choose an assignment to check teacher availability.",
                               bg=WHITE, fg=MUTED, font=("Segoe UI", 9),
                               justify="left", anchor="w", wraplength=550)
        clash_label.pack(fill="x")

        def get_selected_assignment():
            return pair_by_label.get(assignment_var.get().strip(), ("", ""))

        def check_status(_event=None):
            teacher, subject = get_selected_assignment()
            subject_var.set(subject)
            if not teacher or not subject:
                clash_label.configure(text="No teacher/subject assignments found for this grade. Add an assignment to Requirements, then reopen the editor.",
                                      bg=AMBER_PALE, fg=AMBER)
                clash_card.configure(bg=AMBER_PALE, highlightbackground=AMBER)
                return

            clash = self._teacher_clash(teacher, class_name, row_number, col_number)
            if clash:
                other_class, other_subject = clash
                clash_label.configure(
                    text=f"CLASH DETECTED · {teacher} is teaching {other_subject or 'a lesson'} in Class {other_class} at this time. Choose another assignment.",
                    bg=RED_PALE, fg=RED)
                clash_card.configure(bg=RED_PALE, highlightbackground=RED)
            else:
                clash_label.configure(
                    text=f"NO CLASH · {teacher} is available for {subject} during this period.",
                    bg=GREEN_PALE, fg=GREEN)
                clash_card.configure(bg=GREEN_PALE, highlightbackground=GREEN)

        def show_availability_details():
            """Show a populated, searchable view of teachers for this grade and slot."""
            # Refresh assignments at the moment the window opens. This prevents
            # stale/empty lists if the class timetable was edited during this session.
            detail_pairs = self._teacher_subject_pairs_for_grade(class_name)
            seen_pairs = {(item["teacher"].casefold(), item["subject"].casefold())
                          for item in detail_pairs}
            for item in pairs:
                key = (item["teacher"].casefold(), item["subject"].casefold())
                if key not in seen_pairs:
                    detail_pairs.append(dict(item))
                    seen_pairs.add(key)
            detail_pairs.sort(key=lambda item: (item["teacher"].casefold(),
                                                item["subject"].casefold()))

            detail = tk.Toplevel(dialog)
            detail.title(f"Teacher availability · Grade {grade}")
            detail.geometry("760x500")
            detail.minsize(640, 380)
            detail.configure(bg=BG)
            detail.transient(dialog)

            # Set an explicit style: relying on the platform's default Treeview
            # theme can make the row text hard to see on some Windows themes.
            detail_style = ttk.Style(detail)
            detail_style.configure(
                "Availability.Treeview", background=WHITE, foreground=TEXT,
                fieldbackground=WHITE, rowheight=30, borderwidth=0,
                font=("Segoe UI", 9))
            detail_style.configure(
                "Availability.Treeview.Heading", background="#F7F7F5",
                foreground=TEXT, font=("Segoe UI", 9, "bold"), padding=8)
            detail_style.map(
                "Availability.Treeview",
                background=[("selected", "#E3EEFA")],
                foreground=[("selected", TEXT)])

            outer_detail = tk.Frame(detail, bg=BG, padx=18, pady=15)
            outer_detail.pack(fill="both", expand=True)
            heading = tk.Frame(outer_detail, bg=BG)
            heading.pack(fill="x", pady=(0, 10))
            tk.Label(heading, text="Teacher availability", bg=BG, fg=TEXT,
                     font=("Segoe UI", 15, "bold")).pack(anchor="w")
            tk.Label(heading, text=f"{class_name}  ·  {day}  ·  {period}", bg=BG,
                     fg=MUTED, font=("Segoe UI", 9)).pack(anchor="w", pady=(3, 0))

            summary_var = tk.StringVar(value="Loading teacher assignments…")
            tk.Label(outer_detail, textvariable=summary_var, bg=BG, fg=MUTED,
                     font=("Segoe UI", 9), anchor="w").pack(fill="x", pady=(0, 8))

            filter_row = tk.Frame(outer_detail, bg=BG)
            filter_row.pack(fill="x", pady=(0, 8))
            tk.Label(filter_row, text="SEARCH", bg=BG, fg=MUTED,
                     font=("Segoe UI", 8, "bold")).pack(side="left", padx=(0, 8))
            filter_var = tk.StringVar()
            filter_entry = ttk.Entry(filter_row, textvariable=filter_var)
            filter_entry.pack(side="left", fill="x", expand=True)
            filter_entry.configure(font=("Segoe UI", 9))

            table_frame = tk.Frame(outer_detail, bg=BG)
            table_frame.pack(fill="both", expand=True)
            columns = ("teacher", "subject", "availability")
            table = ttk.Treeview(
                table_frame, columns=columns, show="headings", height=12,
                style="Availability.Treeview")
            table.heading("teacher", text="Teacher")
            table.heading("subject", text="Subject")
            table.heading("availability", text="Status")
            table.column("teacher", width=230, minwidth=150, anchor="w", stretch=True)
            table.column("subject", width=180, minwidth=110, anchor="w", stretch=True)
            table.column("availability", width=220, minwidth=160, anchor="w", stretch=True)
            yscroll = ttk.Scrollbar(table_frame, orient="vertical", command=table.yview)
            xscroll = ttk.Scrollbar(table_frame, orient="horizontal", command=table.xview)
            table.configure(yscrollcommand=yscroll.set, xscrollcommand=xscroll.set)
            table.grid(row=0, column=0, sticky="nsew")
            yscroll.grid(row=0, column=1, sticky="ns")
            xscroll.grid(row=1, column=0, sticky="ew")
            table_frame.rowconfigure(0, weight=1)
            table_frame.columnconfigure(0, weight=1)
            table.tag_configure("available", foreground=GREEN)
            table.tag_configure("busy", foreground=RED)

            rows_to_show = []
            for item in detail_pairs:
                teacher = item["teacher"]
                subject = item["subject"]
                clash = self._teacher_clash(teacher, class_name, row_number, col_number)
                if clash:
                    status = f"Busy in Class {clash[0]}"
                    tag = "busy"
                else:
                    subject_conflict = self._teacher_subject_conflict(teacher, subject)
                    if subject_conflict:
                        status = f"Subject conflict: {subject_conflict}"
                        tag = "busy"
                    else:
                        status = "Available"
                        tag = "available"
                rows_to_show.append({
                    "teacher": teacher, "subject": subject,
                    "status": status, "tag": tag})

            empty_label = tk.Label(
                table_frame,
                text=(f"No teacher assignments were found for Grade {grade}.\n"
                      "Check the Requirements worksheet and reopen this window."),
                bg=WHITE, fg=MUTED, font=("Segoe UI", 10), justify="center")

            def refresh_rows(*_args):
                query = filter_var.get().strip().casefold()
                for row_id in table.get_children():
                    table.delete(row_id)
                visible = 0
                available = 0
                busy = 0
                for index, item in enumerate(rows_to_show):
                    searchable = f"{item['teacher']} {item['subject']} {item['status']}".casefold()
                    if query and query not in searchable:
                        continue
                    table.insert("", "end", iid=f"teacher_{index}",
                                 values=(item["teacher"], item["subject"], item["status"]),
                                 tags=(item["tag"],))
                    visible += 1
                    if item["tag"] == "available":
                        available += 1
                    else:
                        busy += 1
                summary_var.set(
                    f"{visible} assignments shown  ·  {available} available  ·  {busy} busy"
                    if rows_to_show else
                    f"No assignments found for Grade {grade}.")
                if not rows_to_show:
                    table.grid_remove()
                    yscroll.grid_remove()
                    xscroll.grid_remove()
                    empty_label.grid(row=0, column=0, sticky="nsew", padx=10, pady=22)
                else:
                    empty_label.grid_remove()
                    table.grid()
                    yscroll.grid()
                    xscroll.grid()
                if rows_to_show and visible == 0:
                    summary_var.set("No assignments match your search. Clear the search field to show all teachers.")

            filter_var.trace_add("write", refresh_rows)
            refresh_rows()

            footer = tk.Frame(outer_detail, bg=BG)
            footer.pack(fill="x", pady=(10, 0))
            tk.Button(footer, text="Close", command=detail.destroy,
                      bg="#F1F1EF", fg=TEXT, relief="flat", padx=16, pady=7,
                      cursor="hand2").pack(side="right")

            detail.update_idletasks()
            detail.lift()
            detail.focus_force()

        tk.Button(availability_row, text="View details", command=show_availability_details,
                  bg=PURPLE_PALE, fg=PURPLE, activebackground="#E3EEFA",
                  activeforeground=PURPLE_DARK, relief="flat", bd=0,
                  padx=10, pady=5, font=("Segoe UI", 8, "bold"),
                  cursor="hand2").pack(side="right", padx=(10, 0))

        def update_assignment_options():
            if not pairs:
                teacher_hint.configure(text=f"No assignments found for Grade {grade}.", fg=AMBER)
                return
            available_count = 0
            busy_count = 0
            for item in pairs:
                clash = self._teacher_clash(item["teacher"], class_name, row_number, col_number)
                if clash:
                    busy_count += 1
                else:
                    available_count += 1
            teacher_hint.configure(
                text=f"Grade {grade} · {len(pair_labels)} assignments · {available_count} available · {busy_count} busy",
                fg=GREEN if available_count else RED)

        # Start with the exact existing assignment when possible. Otherwise,
        # choose a free assignment in this grade before falling back to the first.
        matching_old = next((label for label, pair in pair_by_label.items()
                             if pair[0].casefold() == old_teacher.casefold()
                             and pair[1].casefold() == old_subject.casefold()), None)
        if matching_old:
            initial_label = matching_old
        else:
            available_labels = [label for label, pair in pair_by_label.items()
                                if self._teacher_clash(pair[0], class_name,
                                                       row_number, col_number) is None]
            initial_label = available_labels[0] if available_labels else (pair_labels[0] if pair_labels else "")
        assignment_var.set(initial_label)
        assignment_box.bind("<<ComboboxSelected>>", check_status)
        update_assignment_options()
        check_status()

        buttons = tk.Frame(body, bg=BG)
        buttons.pack(fill="x", side="bottom", pady=(4, 0))

        def save_assignment():
            teacher, subject = get_selected_assignment()
            if not teacher or not subject:
                messagebox.showwarning(
                    "Choose an assignment",
                    f"No available teacher/subject pairs are configured for Grade {grade}. Add the assignment to Requirements first.",
                    parent=dialog)
                assignment_box.focus_set()
                return
            clash = self._teacher_clash(teacher, class_name, row_number, col_number)
            if clash:
                other_class, other_subject = clash
                messagebox.showerror(
                    "Teacher clash",
                    f"{teacher} is already teaching {other_subject or 'a lesson'} in Class {other_class} during {day}, {period}.",
                    parent=dialog)
                check_status()
                return
            try:
                self._apply_lesson_change(class_name, row_number, col_number, subject, teacher)
            except ValueError as error:
                messagebox.showerror("Teacher clash", str(error), parent=dialog)
                check_status()
                return
            dialog.destroy()

        tk.Button(buttons, text="Cancel", command=dialog.destroy, bg="#F1F1EF", fg=TEXT,
                  relief="flat", padx=15, pady=8, cursor="hand2").pack(side="right")
        tk.Button(buttons, text="Save lesson", command=save_assignment, bg=PURPLE,
                  fg=WHITE, activebackground=PURPLE_DARK, relief="flat", padx=18,
                  pady=8, font=("Segoe UI", 9, "bold"), cursor="hand2").pack(side="right", padx=(0, 8))
        if old_subject:
            tk.Button(buttons, text="Clear period",
                      command=lambda: self._clear_from_dialog(dialog, class_name, row_number, col_number),
                      bg=RED_PALE, fg=RED, relief="flat", padx=14, pady=8,
                      cursor="hand2").pack(side="left")
        dialog.bind("<Return>", lambda _event: save_assignment())
        assignment_box.focus_set()

    # Confirm and apply a request to clear a selected lesson.
    def _clear_from_dialog(self, dialog, class_name, row_number, col_number):
        if not messagebox.askyesno("Clear period", "Remove this lesson? The linked teacher sheet and Requirements count will be updated.", parent=dialog):
            return
        self._apply_lesson_change(class_name, row_number, col_number, "", "")
        dialog.destroy()

    # Clear the currently selected period after validating linked timetable data.
    def clear_selected_period(self):
        sheet = self._current_sheet()
        if not self._is_class_sheet(sheet.title):
            messagebox.showinfo("Choose a class sheet", "Select a class timetable before clearing a period.", parent=self.window)
            return
        if not self.selected_cell:
            messagebox.showinfo("Select a period", "Click a period card first, then choose Clear selected period.", parent=self.window)
            return
        row_number, col_number = self.selected_cell
        if row_number < 3 or col_number < 2:
            messagebox.showinfo("Select a period", "Select a lesson card, not a heading.", parent=self.window)
            return
        if not sheet.cell(row_number, col_number).value:
            messagebox.showinfo("Empty period", "This period is already empty.", parent=self.window)
            return
        if messagebox.askyesno("Clear this period?", "Remove this lesson? The teacher sheet and Requirements count will also be updated.", parent=self.window):
            self._apply_lesson_change(sheet.title, row_number, col_number, "", "")

    # Start editing a selected cell when the user double-clicks it.
    def begin_cell_edit(self, event):
        if self.active_entry is not None or self.tree is None:
            return
        if self.tree.identify_region(event.x, event.y) != "cell":
            return
        col_id = self.tree.identify_column(event.x)
        if col_id == "#0":
            return
        item_id = self.tree.identify_row(event.y)
        if not item_id:
            return
        row, col = int(item_id), int(col_id[1:])
        sheet = self._current_sheet()
        cell = sheet.cell(row, col)
        if isinstance(cell, MergedCell):
            return
        box = self.tree.bbox(item_id, col_id)
        if not box:
            return
        x, y, width, height = box
        entry = ttk.Entry(self.tree, font=("Segoe UI", 9))
        entry.place(x=x, y=y, width=width, height=height)
        entry.insert(0, self._display_value(cell.value))
        entry.selection_range(0, tk.END)
        self.active_entry = entry
        completed = {"done": False}

        def commit(show_error=True):
            if completed["done"]:
                return True
            raw = entry.get().strip()
            if sheet.title == "Requirements" and row >= 2 and col == 5:
                try:
                    number = int(raw)
                    if number < 1 or number > 99:
                        raise ValueError
                except ValueError:
                    if show_error:
                        messagebox.showerror("Invalid period count", "Enter a whole number from 1 to 99. To remove an assignment, clear its lessons from the class timetable.", parent=self.window)
                        entry.focus_set()
                        return False
                    return False
                new_value = number
            elif raw == "":
                new_value = None
            elif isinstance(cell.value, int) and not isinstance(cell.value, bool):
                try:
                    new_value = int(raw)
                except ValueError:
                    new_value = raw
            elif isinstance(cell.value, float):
                try:
                    new_value = float(raw)
                except ValueError:
                    new_value = raw
            else:
                new_value = raw
            if cell.value != new_value:
                cell.value = new_value
                self.dirty = True
            completed["done"] = True
            self.active_entry = None
            self.active_commit = None
            entry.destroy()
            self.load_sheet()
            return True

        def cancel():
            if completed["done"]:
                return True
            completed["done"] = True
            self.active_entry = None
            self.active_commit = None
            entry.destroy()
            return True

        self.active_commit = commit
        entry.bind("<Return>", lambda _event: (commit(), "break")[1])
        entry.bind("<Escape>", lambda _event: (cancel(), "break")[1])
        entry.bind("<FocusOut>", lambda _event: commit())
        entry.focus_set()

    # Remember which table cell the user selected.
    def record_selected_cell(self, event):
        if self.tree is None or self.tree.identify_region(event.x, event.y) != "cell":
            return
        item = self.tree.identify_row(event.y)
        col = self.tree.identify_column(event.x)
        if item and col != "#0":
            try:
                self.selected_cell = (int(item), int(col[1:]))
            except (ValueError, IndexError):
                self.selected_cell = None

    # Write all pending edits to the Excel workbook and refresh related app state.
    def save_changes(self):
        if self.active_commit is not None and not self.active_commit():
            return False
        # Validate Requirements and class sheets so manual edits cannot assign
        # one teacher to two different subjects.
        teacher_subjects = {}

        def check_teacher_subject(teacher, subject):
            if not teacher or not subject:
                return True
            key = teacher.strip().casefold()
            previous = teacher_subjects.get(key)
            if previous and previous.casefold() != subject.strip().casefold():
                messagebox.showerror(
                    "Teacher subject rule",
                    f"{teacher} is assigned to both {previous} and {subject}. Each teacher can teach only one subject.\n\nFix the assignments before saving.",
                    parent=self.window,
                )
                return False
            teacher_subjects[key] = subject.strip()
            return True

        if "Requirements" in self.workbook.sheetnames:
            req_sheet = self.workbook["Requirements"]
            for row in range(2, req_sheet.max_row + 1):
                subject = self._display_value(req_sheet.cell(row, 3).value).strip()
                teacher = self._display_value(req_sheet.cell(row, 4).value).strip()
                if not check_teacher_subject(teacher, subject):
                    return False

        for name in self.workbook.sheetnames:
            if self._is_class_sheet(name):
                sheet = self.workbook[name]
                for row in range(3, sheet.max_row + 1):
                    for col in range(2, sheet.max_column + 1):
                        subject, teacher = self._split_lesson(self._display_value(sheet.cell(row, col).value))
                        if not check_teacher_subject(teacher, subject):
                            return False
        try:
            self.workbook.save(self.path)
        except PermissionError:
            messagebox.showerror("Workbook is in use", "Close the workbook in Microsoft Excel, then save again.", parent=self.window)
            return False
        except Exception as error:
            messagebox.showerror("Could not save workbook", str(error), parent=self.window)
            return False
        self.dirty = False
        self.status_var.set("Saved successfully · " + os.path.basename(self.path))
        if self.on_saved:
            self.on_saved(self.path)
        messagebox.showinfo("Workbook saved", "Your Excel workbook changes have been saved.", parent=self.window)
        return True

    # Close the editor safely, handling unsaved changes as needed.
    def close_window(self):
        if self.active_commit is not None and not self.active_commit():
            return
        if self.dirty:
            choice = messagebox.askyesnocancel("Unsaved Excel changes", "Save your changes before closing?", parent=self.window)
            if choice is None:
                return
            if choice and not self.save_changes():
                return
        try:
            self.workbook.close()
        except Exception:
            pass
        self.window.destroy()



# Replace the subject in a class timetable only after checking all slots for teacher clashes.
def replace_class_timetable_subject(workbook_path, class_name, subject):
    """Replace every period in one class with a grade subject, without teacher clashes.

    All slots are checked before modifying the workbook. If no teacher can cover even
    one slot, the whole operation is cancelled rather than leaving a partial timetable.
    """
    workbook_path = os.path.abspath(workbook_path)
    if not os.path.isfile(workbook_path):
        raise ValueError("The selected Excel workbook could not be found.")

    workbook = load_workbook(workbook_path)
    editor = ExcelTimetableEditor.__new__(ExcelTimetableEditor)
    editor.path = workbook_path
    editor.workbook = workbook
    editor._refresh_teacher_sheet_map()

    try:
        if not editor._is_class_sheet(class_name) or class_name not in workbook.sheetnames:
            raise ValueError("Choose a valid class timetable, such as 1-A or 7-C.")
        class_sheet = workbook[class_name]
        grade, division = class_name.split("-", 1)
        selected_subject = str(subject or "").strip()
        if not selected_subject:
            raise ValueError("Choose a subject first.")

        grade_subjects = editor._subjects_for_grade(class_name)
        subject_lookup = {item.casefold(): item for item in grade_subjects}
        if selected_subject.casefold() not in subject_lookup:
            raise ValueError(
                f"{selected_subject} is not listed as a subject available for Grade {grade}. "
                "Choose a subject from the grade's dropdown."
            )
        selected_subject = subject_lookup[selected_subject.casefold()]

        day_rows = [row for row in range(3, class_sheet.max_row + 1)
                    if editor._display_value(class_sheet.cell(row, 1).value).strip()]
        period_cols = list(range(2, class_sheet.max_column + 1))
        slots = [(row, col) for row in day_rows for col in period_cols]
        if not slots:
            raise ValueError("The selected class timetable has no day/period cells.")

        grade_teachers = editor._teachers_for_subject_in_grade(class_name, selected_subject)
        if not grade_teachers:
            raise ValueError(
                f"No teachers are assigned to {selected_subject} in Grade {grade}. "
                "Add a teacher assignment for this subject before replacing the timetable."
            )

        # Work out what subjects each candidate would still teach after the old
        # timetable and Requirements rows for this class are removed.
        def other_subjects_after_replacement(teacher):
            found = set()
            teacher_key = teacher.strip().casefold()
            req = workbook["Requirements"] if "Requirements" in workbook.sheetnames else None
            if req is not None:
                for req_row in range(2, req.max_row + 1):
                    req_grade = editor._display_value(req.cell(req_row, 1).value).strip()
                    req_division = editor._display_value(req.cell(req_row, 2).value).strip().upper()
                    req_subject = editor._display_value(req.cell(req_row, 3).value).strip()
                    req_teacher = editor._display_value(req.cell(req_row, 4).value).strip().casefold()
                    if req_grade == grade and req_division == division.upper():
                        continue
                    if req_teacher == teacher_key and req_subject:
                        found.add(req_subject)

            # The class sheets are authoritative too; omit the selected class
            # because every one of its lessons is being replaced.
            for sheet_name in workbook.sheetnames:
                if not editor._is_class_sheet(sheet_name) or sheet_name == class_name:
                    continue
                other_sheet = workbook[sheet_name]
                for r in range(3, other_sheet.max_row + 1):
                    for c in range(2, other_sheet.max_column + 1):
                        other_subject, other_teacher = editor._split_lesson(
                            editor._display_value(other_sheet.cell(r, c).value))
                        if other_teacher.casefold() == teacher_key and other_subject:
                            found.add(other_subject)
            return found

        eligible_teachers = []
        rejected_for_subject_rule = {}
        for teacher in grade_teachers:
            incompatible = sorted(
                (existing for existing in other_subjects_after_replacement(teacher)
                 if existing.casefold() != selected_subject.casefold()),
                key=str.casefold,
            )
            if incompatible:
                rejected_for_subject_rule[teacher] = incompatible[0]
            else:
                eligible_teachers.append(teacher)

        if not eligible_teachers:
            details = "; ".join(
                f"{teacher} is assigned to {other_subject}"
                for teacher, other_subject in sorted(rejected_for_subject_rule.items(), key=lambda item: item[0].casefold())
            )
            raise ValueError(
                f"There are no eligible teachers for {selected_subject} in Grade {grade}. "
                "The school rule allows each teacher to teach only one subject. " + details
            )

        planned_load = {teacher: 0 for teacher in eligible_teachers}
        plan = []
        blocked_slots = []
        for row, col in slots:
            old_subject, old_teacher = editor._split_lesson(
                editor._display_value(class_sheet.cell(row, col).value))
            available = []
            busy_reasons = []
            for teacher in eligible_teachers:
                clash = editor._teacher_clash(teacher, class_name, row, col)
                if clash:
                    busy_reasons.append(f"{teacher} is busy in {clash[0]}")
                else:
                    available.append(teacher)
            if not available:
                day = editor._display_value(class_sheet.cell(row, 1).value)
                period = editor._display_value(class_sheet.cell(2, col).value) or f"Period {col - 1}"
                blocked_slots.append(f"{day}, {period}: " + (", ".join(busy_reasons) or "no eligible teacher is free"))
                continue

            # Keep the current teacher when the cell already has the desired
            # subject. For changed cells, distribute assignments across available
            # teachers instead of assigning every slot to one person.
            available.sort(key=lambda teacher: (
                0 if old_subject.casefold() == selected_subject.casefold()
                and old_teacher.casefold() == teacher.casefold() else 1,
                planned_load[teacher],
                teacher.casefold(),
            ))
            chosen_teacher = available[0]
            planned_load[chosen_teacher] += 1
            plan.append((row, col, chosen_teacher))

        if blocked_slots:
            shown = "\n".join("• " + item for item in blocked_slots[:8])
            more = "" if len(blocked_slots) <= 8 else f"\n…and {len(blocked_slots) - 8} more period(s)."
            raise ValueError(
                f"Cannot replace the entire timetable for Class {class_name} with {selected_subject}. "
                f"No changes were made because a matching teacher is unavailable in {len(blocked_slots)} period(s):\n\n"
                f"{shown}{more}\n\nChoose another subject or adjust other classes' teacher assignments first."
            )

        # Clear all class-linked entries from teacher sheets first, then write
        # the new lesson grid. This also repairs any stale teacher-cell entries.
        for sheet_name in workbook.sheetnames:
            if not editor._is_teacher_sheet(sheet_name):
                continue
            teacher_sheet = workbook[sheet_name]
            for row in range(3, teacher_sheet.max_row + 1):
                for col in range(2, teacher_sheet.max_column + 1):
                    assigned_class, _assigned_subject = editor._split_lesson(
                        editor._display_value(teacher_sheet.cell(row, col).value))
                    if assigned_class == class_name:
                        teacher_sheet.cell(row, col).value = None

        # Clear the selected class's old requirement rows, then rebuild accurate
        # weekly counts from the teacher assignments chosen for each period.
        if "Requirements" in workbook.sheetnames:
            req = workbook["Requirements"]
            for req_row in range(req.max_row, 1, -1):
                req_grade = editor._display_value(req.cell(req_row, 1).value).strip()
                req_division = editor._display_value(req.cell(req_row, 2).value).strip().upper()
                if req_grade == grade and req_division == division.upper():
                    req.delete_rows(req_row, 1)

        counts = {}
        for row, col, teacher in plan:
            class_sheet.cell(row, col).value = selected_subject + "\n" + teacher
            teacher_sheet_name = editor._populate_teacher_sheet_if_missing(teacher)
            workbook[teacher_sheet_name].cell(row, col).value = class_name + "\n" + selected_subject
            counts[teacher] = counts.get(teacher, 0) + 1

        if "Requirements" in workbook.sheetnames:
            req = workbook["Requirements"]
            for teacher in sorted(counts, key=str.casefold):
                req.append([int(grade), division, selected_subject, teacher, counts[teacher]])

        # Remove adjustment-history rows for this class because its old subject
        # allocations no longer describe the timetable after bulk replacement.
        if "Repeat Adjustments" in workbook.sheetnames:
            adjustments = workbook["Repeat Adjustments"]
            for adjust_row in range(adjustments.max_row, 1, -1):
                if editor._display_value(adjustments.cell(adjust_row, 1).value).strip() == class_name:
                    adjustments.delete_rows(adjust_row, 1)
            if "Summary" in workbook.sheetnames:
                workbook["Summary"].cell(14, 2).value = max(0, adjustments.max_row - 1)

        # Refresh the teacher-subject column in Summary using the updated inputs.
        if "Summary" in workbook.sheetnames:
            summary = workbook["Summary"]
            subjects_by_teacher = {}
            req = workbook["Requirements"] if "Requirements" in workbook.sheetnames else None
            if req is not None:
                for req_row in range(2, req.max_row + 1):
                    subject_value = editor._display_value(req.cell(req_row, 3).value).strip()
                    teacher_value = editor._display_value(req.cell(req_row, 4).value).strip()
                    if subject_value and teacher_value:
                        subjects_by_teacher.setdefault(teacher_value.casefold(), set()).add(subject_value)
            for sheet_name in workbook.sheetnames:
                if not editor._is_class_sheet(sheet_name):
                    continue
                other_sheet = workbook[sheet_name]
                for r in range(3, other_sheet.max_row + 1):
                    for c in range(2, other_sheet.max_column + 1):
                        subject_value, teacher_value = editor._split_lesson(
                            editor._display_value(other_sheet.cell(r, c).value))
                        if subject_value and teacher_value:
                            subjects_by_teacher.setdefault(teacher_value.casefold(), set()).add(subject_value)
            for summary_row in range(2, summary.max_row + 1):
                teacher_value = editor._display_value(summary.cell(summary_row, 4).value).strip()
                if teacher_value:
                    teacher_subjects = subjects_by_teacher.get(teacher_value.casefold(), set())
                    summary.cell(summary_row, 6).value = (
                        sorted(teacher_subjects, key=str.casefold)[0] if len(teacher_subjects) == 1 else ""
                    )

        # Persist only after all slots have an available, subject-eligible teacher.
        try:
            workbook.save(workbook_path)
        except PermissionError as error:
            raise ValueError("Close the Excel workbook in Microsoft Excel, then try again.") from error

        return {
            "class": class_name,
            "subject": selected_subject,
            "periods": len(plan),
            "teachers": sorted(counts, key=str.casefold),
            "teacher_counts": counts,
        }
    finally:
        workbook.close()


# Count scheduled periods for the selected teacher and subject in each matching division.
def _teacher_subject_division_counts(workbook, grade, subject, teacher):
    """Return only divisions where this teacher actually teaches this subject,
    with the scheduled weekly period count for each division.
    """
    grade = str(grade).strip()
    subject_key = str(subject).strip().casefold()
    teacher_key = str(teacher).strip().casefold()
    counts = {}
    if not subject_key or not teacher_key:
        return counts

    for class_name in workbook.sheetnames:
        if not CLASS_RE.match(class_name) or class_name.split("-", 1)[0] != grade:
            continue
        sheet = workbook[class_name]
        count = 0
        for row in range(3, sheet.max_row + 1):
            if not ExcelTimetableEditor._display_value(sheet.cell(row, 1).value).strip():
                continue
            for col in range(2, sheet.max_column + 1):
                cell_subject, cell_teacher = ExcelTimetableEditor._split_lesson(
                    ExcelTimetableEditor._display_value(sheet.cell(row, col).value)
                )
                if (cell_subject.casefold() == subject_key and
                        cell_teacher.casefold() == teacher_key):
                    count += 1
        if count:
            division = class_name.split("-", 1)[1]
            counts[division] = count
    return dict(sorted(counts.items(), key=lambda item: item[0].casefold()))


# Build the subject and teacher choices available for one grade.
def _split_subject_data(workbook, grade):
    """Return subjects and teachers actually assigned to a subject in one grade."""
    grade = str(grade).strip()
    subjects = {}
    teachers_by_subject = {}
    subject_counts = {}

    # Requirements supplies initial choices; class timetable cells confirm the
    # assignments that are actually scheduled in the workbook.
    if "Requirements" in workbook.sheetnames:
        req = workbook["Requirements"]
        for row in range(2, req.max_row + 1):
            row_grade = ExcelTimetableEditor._display_value(req.cell(row, 1).value).strip()
            subject = ExcelTimetableEditor._display_value(req.cell(row, 3).value).strip()
            teacher = ExcelTimetableEditor._display_value(req.cell(row, 4).value).strip()
            if row_grade == grade and subject and teacher:
                key = subject.casefold()
                subjects.setdefault(key, subject)
                teachers_by_subject.setdefault(key, {})[teacher.casefold()] = teacher

    for name in workbook.sheetnames:
        if not CLASS_RE.match(name) or name.split("-", 1)[0] != grade:
            continue
        sheet = workbook[name]
        for row in range(3, sheet.max_row + 1):
            if not ExcelTimetableEditor._display_value(sheet.cell(row, 1).value).strip():
                continue
            for col in range(2, sheet.max_column + 1):
                subject, teacher = ExcelTimetableEditor._split_lesson(
                    ExcelTimetableEditor._display_value(sheet.cell(row, col).value)
                )
                if subject and teacher:
                    key = subject.casefold()
                    subjects.setdefault(key, subject)
                    teachers_by_subject.setdefault(key, {})[teacher.casefold()] = teacher
                    subject_counts[(name, key, teacher.casefold())] = (
                        subject_counts.get((name, key, teacher.casefold()), 0) + 1
                    )

    return (
        sorted(subjects.values(), key=str.casefold),
        teachers_by_subject,
        subject_counts,
    )


# Map teacher display names to the workbook worksheet names used for their timetables.
def _teacher_sheet_lookup(workbook):
    """Map teacher display names to the combined teacher worksheet names."""
    result = {}
    if "Summary" in workbook.sheetnames:
        summary = workbook["Summary"]
        for row in range(2, summary.max_row + 1):
            teacher = ExcelTimetableEditor._display_value(summary.cell(row, 4).value).strip()
            tab = ExcelTimetableEditor._display_value(summary.cell(row, 5).value).strip()
            if teacher and tab in workbook.sheetnames:
                result[teacher.casefold()] = (teacher, tab)

    for tab in workbook.sheetnames:
        if not TEACHER_RE.match(tab):
            continue
        title = ExcelTimetableEditor._display_value(workbook[tab].cell(1, 1).value).strip()
        if title.casefold().startswith("teacher:"):
            teacher = title.split(":", 1)[1].strip()
            if teacher:
                result.setdefault(teacher.casefold(), (teacher, tab))
    return result


# Create a teacher timetable worksheet using the workbook’s existing layout and styles.
def _create_teacher_sheet(workbook, teacher, teacher_lookup):
    """Create a new teacher sheet using the workbook's existing timetable style."""
    existing = teacher_lookup.get(teacher.casefold())
    if existing and existing[1] in workbook.sheetnames:
        return existing[1]

    numbers = []
    for name in workbook.sheetnames:
        if TEACHER_RE.match(name):
            try:
                numbers.append(int(name[1:]))
            except ValueError:
                pass
    sheet_name = "T" + str(max(numbers + [0]) + 1).zfill(3)

    template = next((workbook[name] for name in workbook.sheetnames if TEACHER_RE.match(name)), None)
    if template is not None:
        sheet = workbook.copy_worksheet(template)
        sheet.title = sheet_name
        # Keep the template's sizing and formatting, but clear the copied lessons.
        for row in range(3, sheet.max_row + 1):
            for col in range(2, sheet.max_column + 1):
                sheet.cell(row, col).value = None
        sheet.cell(1, 1).value = "Teacher: " + teacher
    else:
        sheet = workbook.create_sheet(sheet_name)
        sample_class = next((workbook[name] for name in workbook.sheetnames if CLASS_RE.match(name)), None)
        period_count = max(1, sample_class.max_column - 1) if sample_class is not None else 7
        day_names = ([ExcelTimetableEditor._display_value(sample_class.cell(row, 1).value)
                      for row in range(3, sample_class.max_row + 1)
                      if ExcelTimetableEditor._display_value(sample_class.cell(row, 1).value).strip()]
                     if sample_class is not None else
                     ["Monday", "Tuesday", "Wednesday", "Thursday", "Friday"])
        sheet.merge_cells(start_row=1, start_column=1, end_row=1, end_column=period_count + 1)
        sheet.cell(1, 1, "Teacher: " + teacher)
        sheet.cell(2, 1, "Day")
        for period in range(period_count):
            sheet.cell(2, period + 2, "Period " + str(period + 1))
        for index, day in enumerate(day_names):
            sheet.cell(index + 3, 1, day)
        sheet.column_dimensions["A"].width = 15
        for col in range(2, period_count + 2):
            sheet.column_dimensions[get_column_letter(col)].width = 18

    teacher_lookup[teacher.casefold()] = (teacher, sheet_name)
    return sheet_name


# Check class and teacher worksheets for an existing lesson at a candidate time slot.
def _teacher_is_free_in_workbook(workbook, teacher, class_name, row, col, teacher_lookup):
    """Check both the teacher's sheet and class sheets for a time-slot clash."""
    lookup = teacher_lookup.get(teacher.casefold())
    if lookup and lookup[1] in workbook.sheetnames:
        value = ExcelTimetableEditor._display_value(workbook[lookup[1]].cell(row, col).value)
        if value:
            assigned_class, _assigned_subject = ExcelTimetableEditor._split_lesson(value)
            if assigned_class and assigned_class != class_name:
                return False, assigned_class

    for name in workbook.sheetnames:
        if not CLASS_RE.match(name) or name == class_name:
            continue
        sheet = workbook[name]
        value = ExcelTimetableEditor._display_value(sheet.cell(row, col).value)
        _subject, assigned_teacher = ExcelTimetableEditor._split_lesson(value)
        if assigned_teacher and assigned_teacher.casefold() == teacher.casefold():
            return False, name
    return True, ""


# Refresh the Summary sheet’s teacher list, sheet IDs, and assigned subjects.
def _refresh_teacher_summary(workbook, teacher_lookup):
    """Refresh the teacher list and single-subject column on the Summary sheet."""
    if "Summary" not in workbook.sheetnames:
        summary = workbook.create_sheet("Summary", 0)
    else:
        summary = workbook["Summary"]

    # Teacher sheets are the durable teacher roster; include teachers even if
    # they currently have no assigned periods left after the split.
    for tab in workbook.sheetnames:
        if TEACHER_RE.match(tab):
            title = ExcelTimetableEditor._display_value(workbook[tab].cell(1, 1).value).strip()
            if title.casefold().startswith("teacher:"):
                teacher = title.split(":", 1)[1].strip()
                if teacher:
                    teacher_lookup[teacher.casefold()] = (teacher, tab)

    subjects_by_teacher = {}
    for name in workbook.sheetnames:
        if not CLASS_RE.match(name):
            continue
        sheet = workbook[name]
        for row in range(3, sheet.max_row + 1):
            for col in range(2, sheet.max_column + 1):
                subject, teacher = ExcelTimetableEditor._split_lesson(
                    ExcelTimetableEditor._display_value(sheet.cell(row, col).value)
                )
                if subject and teacher:
                    subjects_by_teacher.setdefault(teacher.casefold(), set()).add(subject)
                    teacher_lookup.setdefault(teacher.casefold(), (teacher, ""))

    summary.cell(1, 4, "Teacher")
    summary.cell(1, 5, "Sheet")
    summary.cell(1, 6, "Subject")
    for row in range(2, summary.max_row + 1):
        summary.cell(row, 4).value = None
        summary.cell(row, 5).value = None
        summary.cell(row, 6).value = None

    teachers = sorted(teacher_lookup.values(), key=lambda item: item[0].casefold())
    for row, (teacher, tab) in enumerate(teachers, start=2):
        summary.cell(row, 4, teacher)
        summary.cell(row, 5, tab)
        subjects = subjects_by_teacher.get(teacher.casefold(), set())
        summary.cell(row, 6, sorted(subjects, key=str.casefold)[0] if len(subjects) == 1 else ", ".join(sorted(subjects, key=str.casefold)))

    for row in range(2, summary.max_row + 1):
        if ExcelTimetableEditor._display_value(summary.cell(row, 1).value).strip().casefold() == "number of teachers":
            summary.cell(row, 2, len(teachers))
            break


# Transfer periods to another teacher while checking conflicts and updating all affected sheets.
def split_subject_periods(workbook_path, grade, subject, existing_teacher, new_teacher,
                          periods_per_division=1, division="All divisions"):
    """Transfer subject periods in one chosen division or every division in a grade.

    The requested number is transferred per selected division, but only for
    divisions where the existing teacher currently teaches the subject. The
    transfer is atomic: if any conflict prevents the complete split, nothing
    is saved to the workbook.
    """
    import os
    import tempfile

    grade = str(grade).strip()
    subject = str(subject).strip()
    existing_teacher = str(existing_teacher).strip()
    new_teacher = str(new_teacher).strip()
    try:
        periods_per_division = int(periods_per_division)
    except (TypeError, ValueError):
        raise ValueError("Enter a whole number of periods to transfer per division.")
    if not grade.isdigit() or not 1 <= int(grade) <= 12:
        raise ValueError("Choose a grade between 1 and 12.")
    if not subject or not existing_teacher or not new_teacher:
        raise ValueError("Choose a subject and existing teacher, then enter the new teacher's name.")
    if periods_per_division < 1:
        raise ValueError("At least one period must be transferred per affected division.")
    if existing_teacher.casefold() == new_teacher.casefold():
        raise ValueError("The existing teacher and new teacher must have different names.")

    workbook = load_workbook(workbook_path)
    try:
        if "Requirements" not in workbook.sheetnames:
            raise ValueError("This workbook has no Requirements worksheet. Generate a timetable first.")
        grade_int = int(grade)
        teacher_lookup = _teacher_sheet_lookup(workbook)

        # Normalize the new teacher's name if it already exists in the workbook.
        new_key = new_teacher.casefold()
        if new_key in teacher_lookup:
            new_teacher = teacher_lookup[new_key][0]
        else:
            req = workbook["Requirements"]
            for req_row in range(2, req.max_row + 1):
                req_teacher = ExcelTimetableEditor._display_value(req.cell(req_row, 4).value).strip()
                if req_teacher.casefold() == new_key:
                    new_teacher = req_teacher
                    break

        # Enforce the project-wide rule: each teacher is assigned to one subject.
        assigned_subjects = set()
        req = workbook["Requirements"]
        for req_row in range(2, req.max_row + 1):
            req_teacher = ExcelTimetableEditor._display_value(req.cell(req_row, 4).value).strip()
            req_subject = ExcelTimetableEditor._display_value(req.cell(req_row, 3).value).strip()
            if req_teacher.casefold() == new_teacher.casefold() and req_subject:
                assigned_subjects.add(req_subject)
        for name in workbook.sheetnames:
            if not CLASS_RE.match(name):
                continue
            sheet = workbook[name]
            for row in range(3, sheet.max_row + 1):
                for col in range(2, sheet.max_column + 1):
                    cell_subject, cell_teacher = ExcelTimetableEditor._split_lesson(
                        ExcelTimetableEditor._display_value(sheet.cell(row, col).value)
                    )
                    if cell_teacher.casefold() == new_teacher.casefold() and cell_subject:
                        assigned_subjects.add(cell_subject)
        other_subjects = sorted((item for item in assigned_subjects
                                 if item.casefold() != subject.casefold()), key=str.casefold)
        if other_subjects:
            raise ValueError(
                f"{new_teacher} is already assigned to {', '.join(other_subjects)}. "
                f"Each teacher can teach only one subject, so choose a teacher assigned to {subject} or enter a new name."
            )

        # Canonicalize displayed names and reject teachers that are not already
        # teaching the selected subject in the selected grade.
        subjects, teachers_by_subject, _counts = _split_subject_data(workbook, grade)
        subject_lookup = {item.casefold(): item for item in subjects}
        if subject.casefold() not in subject_lookup:
            raise ValueError(f"{subject} is not currently assigned in Grade {grade}.")
        subject = subject_lookup[subject.casefold()]
        available_old_names = teachers_by_subject.get(subject.casefold(), {})
        if existing_teacher.casefold() not in available_old_names:
            raise ValueError(f"{existing_teacher} is not assigned to {subject} in Grade {grade}.")
        existing_teacher = available_old_names[existing_teacher.casefold()]

        grade_classes = sorted(
            [name for name in workbook.sheetnames
             if CLASS_RE.match(name) and name.split("-", 1)[0] == grade],
            key=str.casefold,
        )
        division = str(division or "All divisions").strip()
        if division.casefold() == "all divisions":
            target_classes = grade_classes[:]
        else:
            target_classes = [
                name for name in grade_classes
                if name.split("-", 1)[1].casefold() == division.casefold()
            ]
            if not target_classes:
                valid_divisions = sorted({name.split("-", 1)[1] for name in grade_classes}, key=str.casefold)
                raise ValueError(
                    f"Division {division} was not found in Grade {grade}. "
                    f"Available divisions: {', '.join(valid_divisions)}."
                )

        slots_by_class = {}
        candidate_slots_by_class = {}
        for class_name in target_classes:
            sheet = workbook[class_name]
            slots = []
            for row in range(3, sheet.max_row + 1):
                if not ExcelTimetableEditor._display_value(sheet.cell(row, 1).value).strip():
                    continue
                for col in range(2, sheet.max_column + 1):
                    cell_subject, cell_teacher = ExcelTimetableEditor._split_lesson(
                        ExcelTimetableEditor._display_value(sheet.cell(row, col).value)
                    )
                    if cell_subject.casefold() == subject.casefold() and cell_teacher.casefold() == existing_teacher.casefold():
                        slots.append((row, col))
            if slots:
                slots_by_class[class_name] = sorted(slots)
                free_slots = []
                for row, col in sorted(slots):
                    is_free, busy_class = _teacher_is_free_in_workbook(
                        workbook, new_teacher, class_name, row, col, teacher_lookup
                    )
                    if is_free:
                        free_slots.append((row, col))
                candidate_slots_by_class[class_name] = free_slots

        if not slots_by_class:
            raise ValueError(
                f"No scheduled {subject} periods taught by {existing_teacher} were found in Grade {grade}."
            )

        short_classes = [
            f"{class_name}: requested {periods_per_division}, only {len(candidate_slots_by_class.get(class_name, []))} clash-free period(s) available"
            for class_name in slots_by_class
            if len(candidate_slots_by_class.get(class_name, [])) < periods_per_division
        ]
        if short_classes:
            raise ValueError(
                "The split cannot be completed without a teacher clash. No changes were made.\n\n" +
                "\n".join(short_classes)
            )

        # Match each requested transfer to a unique day/period. This prevents
        # assigning the new teacher to two different divisions simultaneously.
        left_nodes = []
        ordered_candidates = {}
        for class_name, candidates in candidate_slots_by_class.items():
            count = len(candidates)
            preferred = []
            for index in range(periods_per_division):
                candidate_index = min(count - 1, int((index + 0.5) * count / periods_per_division))
                slot = candidates[candidate_index]
                if slot not in preferred:
                    preferred.append(slot)
            ordered_candidates[class_name] = preferred + [slot for slot in candidates if slot not in preferred]
            for ordinal in range(periods_per_division):
                left_nodes.append((class_name, ordinal))
        left_nodes.sort(key=lambda node: (len(ordered_candidates[node[0]]), node[0], node[1]))

        slot_to_node = {}

        def assign_slot(node, visited):
            class_name, _ordinal = node
            for slot in ordered_candidates[class_name]:
                if slot in visited:
                    continue
                visited.add(slot)
                previous = slot_to_node.get(slot)
                if previous is None or assign_slot(previous, visited):
                    slot_to_node[slot] = node
                    return True
            return False

        for node in left_nodes:
            if not assign_slot(node, set()):
                raise ValueError(
                    "There are not enough distinct periods to share the existing teacher's workload "
                    "without double-booking the new teacher. No changes were made. Try transferring fewer periods per division."
                )

        planned = [(node[0], slot[0], slot[1]) for slot, node in slot_to_node.items()]
        if len(planned) != len(left_nodes):
            raise ValueError("The split could not be planned completely. No changes were made.")

        # Count the subject's scheduled periods before mutation; a successful
        # transfer must preserve these per-division totals exactly.
        before_totals = {}
        for class_name in target_classes:
            sheet = workbook[class_name]
            count = 0
            for row in range(3, sheet.max_row + 1):
                for col in range(2, sheet.max_column + 1):
                    cell_subject, _cell_teacher = ExcelTimetableEditor._split_lesson(
                        ExcelTimetableEditor._display_value(sheet.cell(row, col).value)
                    )
                    if cell_subject.casefold() == subject.casefold():
                        count += 1
            before_totals[class_name] = count

        # Create the teacher sheet only after a full clash-free plan exists.
        new_teacher_sheet = _create_teacher_sheet(workbook, new_teacher, teacher_lookup)
        old_teacher_info = teacher_lookup.get(existing_teacher.casefold())
        old_teacher_sheet = old_teacher_info[1] if old_teacher_info else None
        changed_by_class = {}
        for class_name, row, col in planned:
            class_sheet = workbook[class_name]
            class_sheet.cell(row, col).value = subject + "\n" + new_teacher
            if old_teacher_sheet and old_teacher_sheet in workbook.sheetnames:
                old_cell = workbook[old_teacher_sheet].cell(row, col)
                assigned_class, assigned_subject = ExcelTimetableEditor._split_lesson(
                    ExcelTimetableEditor._display_value(old_cell.value)
                )
                if assigned_class == class_name and assigned_subject.casefold() == subject.casefold():
                    old_cell.value = None
            workbook[new_teacher_sheet].cell(row, col).value = class_name + "\n" + subject
            changed_by_class[class_name] = changed_by_class.get(class_name, 0) + 1

        # Rebuild this grade/subject's requirement rows from the actual class
        # timetables, leaving every other subject and grade untouched.
        req = workbook["Requirements"]
        for req_row in range(req.max_row, 1, -1):
            row_grade = ExcelTimetableEditor._display_value(req.cell(req_row, 1).value).strip()
            row_subject = ExcelTimetableEditor._display_value(req.cell(req_row, 3).value).strip()
            if row_grade == grade and row_subject.casefold() == subject.casefold():
                req.delete_rows(req_row, 1)
        requirement_rows = []
        for class_name in grade_classes:
            sheet = workbook[class_name]
            counts = {}
            for row in range(3, sheet.max_row + 1):
                for col in range(2, sheet.max_column + 1):
                    cell_subject, cell_teacher = ExcelTimetableEditor._split_lesson(
                        ExcelTimetableEditor._display_value(sheet.cell(row, col).value)
                    )
                    if cell_subject.casefold() == subject.casefold() and cell_teacher:
                        counts[cell_teacher] = counts.get(cell_teacher, 0) + 1
            for teacher in sorted(counts, key=str.casefold):
                requirement_rows.append([grade_int, class_name.split("-", 1)[1], subject,
                                         teacher, counts[teacher]])
        for req_row in requirement_rows:
            req.append(req_row)

        # Verify no slot was modified incorrectly and every subject total is stable.
        for class_name in target_classes:
            sheet = workbook[class_name]
            after = 0
            for row in range(3, sheet.max_row + 1):
                for col in range(2, sheet.max_column + 1):
                    cell_subject, _teacher = ExcelTimetableEditor._split_lesson(
                        ExcelTimetableEditor._display_value(sheet.cell(row, col).value)
                    )
                    if cell_subject.casefold() == subject.casefold():
                        after += 1
            if after != before_totals[class_name]:
                raise ValueError(f"Subject totals changed unexpectedly in {class_name}; no workbook was saved.")

        # Rebuild teacher sheet names and summary details, including the new teacher.
        _refresh_teacher_summary(workbook, teacher_lookup)

        # Write to a temporary file in the same directory, then replace the old
        # workbook only after openpyxl has saved a complete file successfully.
        folder = os.path.dirname(os.path.abspath(workbook_path)) or os.getcwd()
        fd, temp_path = tempfile.mkstemp(prefix="schoolflow_split_", suffix=".xlsx", dir=folder)
        os.close(fd)
        try:
            workbook.save(temp_path)
            workbook.close()
            os.replace(temp_path, workbook_path)
        except PermissionError as error:
            try:
                os.remove(temp_path)
            except OSError:
                pass
            raise ValueError("Close the workbook in Microsoft Excel before splitting subject periods.") from error
        except Exception:
            try:
                os.remove(temp_path)
            except OSError:
                pass
            raise

        return {
            "grade": grade,
            "subject": subject,
            "existing_teacher": existing_teacher,
            "new_teacher": new_teacher,
            "divisions": len(slots_by_class),
            "periods_transferred": len(planned),
            "periods_per_division": periods_per_division,
            "division": division,
            "class_counts": changed_by_class,
        }
    finally:
        try:
            workbook.close()
        except Exception:
            pass


# Display the split-period dialog with only relevant teachers and divisions.
def open_split_subject_periods(parent, workbook_path, on_saved=None):
    """Open the split dialog, scoped to divisions where the selected teacher teaches."""
    import os

    if not os.path.isfile(workbook_path):
        messagebox.showerror("Workbook not found", "Generate a timetable before splitting subject periods.", parent=parent)
        return None
    try:
        source_workbook = load_workbook(workbook_path)
    except Exception as error:
        messagebox.showerror("Could not read workbook", str(error), parent=parent)
        return None

    grade_values = sorted(
        {name.split("-", 1)[0] for name in source_workbook.sheetnames if CLASS_RE.match(name)},
        key=lambda value: int(value),
    )
    if not grade_values:
        source_workbook.close()
        messagebox.showerror("No class timetables", "This workbook does not contain class timetable sheets.", parent=parent)
        return None

    window = tk.Toplevel(parent)
    window.title("Split subject periods | SchoolFlow")
    window.configure(bg=BG)
    screen_width, screen_height = window.winfo_screenwidth(), window.winfo_screenheight()
    window.geometry(f"{min(760, max(620, screen_width - 100))}x{min(680, max(600, screen_height - 100))}")
    window.minsize(600, 570)
    window.transient(parent)

    outer = tk.Frame(window, bg=BG, padx=22, pady=20)
    outer.pack(fill="both", expand=True)
    tk.Label(outer, text="Split subject periods", bg=BG, fg=TEXT,
             font=("Segoe UI", 18, "bold")).pack(anchor="w")
    tk.Label(outer,
             text="Transfer lessons from an existing teacher to a new teacher in a division where the existing teacher currently teaches this subject.",
             bg=BG, fg=MUTED, font=("Segoe UI", 9), wraplength=670,
             justify="left").pack(anchor="w", pady=(5, 16))

    card = tk.Frame(outer, bg=WHITE, highlightbackground=BORDER, highlightthickness=1, padx=18, pady=16)
    card.pack(fill="x")
    form = tk.Frame(card, bg=WHITE)
    form.pack(fill="x")
    form.columnconfigure(1, weight=1)

    def label(row, text):
        tk.Label(form, text=text, bg=WHITE, fg=TEXT,
                 font=("Segoe UI", 9, "bold")).grid(row=row, column=0, sticky="w", padx=(0, 15), pady=7)

    grade_var = tk.StringVar(value=grade_values[-1])
    division_var = tk.StringVar(value="")
    subject_var = tk.StringVar()
    existing_var = tk.StringVar()
    new_teacher_var = tk.StringVar()
    transfer_var = tk.StringVar(value="1")

    label(0, "Grade")
    grade_box = ttk.Combobox(form, textvariable=grade_var, values=grade_values, state="readonly", width=32)
    grade_box.grid(row=0, column=1, sticky="ew", pady=7)
    label(1, "Division")
    division_box = ttk.Combobox(form, textvariable=division_var, state="readonly", width=32)
    division_box.grid(row=1, column=1, sticky="ew", pady=7)
    label(2, "Subject")
    subject_box = ttk.Combobox(form, textvariable=subject_var, state="readonly", width=32)
    subject_box.grid(row=2, column=1, sticky="ew", pady=7)
    label(3, "Existing teacher")
    existing_box = ttk.Combobox(form, textvariable=existing_var, state="readonly", width=32)
    existing_box.grid(row=3, column=1, sticky="ew", pady=7)
    label(4, "New teacher name")
    new_entry = ttk.Entry(form, textvariable=new_teacher_var, width=35)
    new_entry.grid(row=4, column=1, sticky="ew", pady=7)
    label(5, "Periods to transfer")
    amount_box = ttk.Spinbox(form, from_=1, to=35, textvariable=transfer_var, width=8)
    amount_box.grid(row=5, column=1, sticky="w", pady=7)
    tk.Label(form, text="in the selected division", bg=WHITE, fg=MUTED,
             font=("Segoe UI", 8)).grid(row=5, column=1, sticky="w", padx=(110, 0), pady=7)

    info_var = tk.StringVar(value="Choose a grade, subject, and existing teacher.")
    info_label = tk.Label(card, textvariable=info_var, bg=PURPLE_PALE, fg=PURPLE,
                          font=("Segoe UI", 9), wraplength=660, justify="left", padx=12, pady=10)
    info_label.pack(fill="x", pady=(14, 0))

    def update_divisions(_event=None):
        """Limit the dropdown to divisions where this teacher has this subject."""
        subject = subject_var.get().strip()
        teacher = existing_var.get().strip()
        matching = _teacher_subject_division_counts(
            source_workbook, grade_var.get(), subject, teacher
        )
        values = list(matching.keys())
        division_box.configure(values=values)
        if division_var.get() not in values:
            division_var.set(values[0] if values else "")
        selected_count = matching.get(division_var.get(), 0)
        amount_box.configure(to=max(1, selected_count))
        try:
            requested = int(transfer_var.get())
        except ValueError:
            requested = 1
        if selected_count and requested > selected_count:
            transfer_var.set(str(selected_count))
        update_info()

    def grade_changed(_event=None):
        update_subjects()

    def update_subjects(_event=None):
        subjects, _teachers_by_subject, _counts = _split_subject_data(source_workbook, grade_var.get())
        subject_box.configure(values=subjects)
        current = subject_var.get()
        if current not in subjects:
            subject_var.set(subjects[0] if subjects else "")
        update_teachers()

    def update_teachers(_event=None):
        _subjects, teachers_by_subject, _counts = _split_subject_data(source_workbook, grade_var.get())
        subject = subject_var.get().casefold()
        teachers = sorted(teachers_by_subject.get(subject, {}).values(), key=str.casefold)
        existing_box.configure(values=teachers)
        if existing_var.get() not in teachers:
            existing_var.set(teachers[0] if teachers else "")
        update_divisions()

    def update_info(_event=None):
        subject = subject_var.get().strip()
        old_teacher = existing_var.get().strip()
        division = division_var.get().strip()
        if not subject or not old_teacher:
            info_var.set("Choose a subject and existing teacher to see matching divisions.")
            return

        matching = _teacher_subject_division_counts(
            source_workbook, grade_var.get(), subject, old_teacher
        )
        if not matching:
            info_var.set(
                f"{old_teacher} has no scheduled {subject} periods in Grade {grade_var.get()}. "
                "Choose another teacher or subject."
            )
            return
        if not division or division not in matching:
            info_var.set(
                f"{old_teacher} teaches {subject} in: " + ", ".join(
                    f"{grade_var.get()}-{name}" for name in matching
                ) + ". Select one of these divisions."
            )
            return

        count = matching[division]
        try:
            requested = int(transfer_var.get())
        except ValueError:
            requested = 0
        if requested < 1:
            info_var.set(
                f"Scope: Grade {grade_var.get()}, Division {division}. "
                f"{old_teacher} teaches {count} {subject} period(s) in this division. "
                "Enter a positive number of periods to transfer."
            )
        elif requested > count:
            info_var.set(
                f"Scope: Grade {grade_var.get()}, Division {division}. "
                f"{old_teacher} teaches {count} {subject} period(s) in this division. "
                f"You can transfer at most {count} period(s)."
            )
        else:
            info_var.set(
                f"Scope: Grade {grade_var.get()}, Division {division}. "
                f"{old_teacher} currently teaches {count} {subject} period(s) in this division. "
                f"This will transfer {requested} period(s) from Grade {grade_var.get()}-{division}. "
                "Only this division is included in the split; other divisions are not changed."
            )

    def division_changed(_event=None):
        matching = _teacher_subject_division_counts(
            source_workbook, grade_var.get(), subject_var.get(), existing_var.get()
        )
        selected_count = matching.get(division_var.get(), 0)
        amount_box.configure(to=max(1, selected_count))
        try:
            requested = int(transfer_var.get())
        except ValueError:
            requested = 1
        if selected_count and requested > selected_count:
            transfer_var.set(str(selected_count))
        update_info()

    grade_box.bind("<<ComboboxSelected>>", grade_changed)
    division_box.bind("<<ComboboxSelected>>", division_changed)
    subject_box.bind("<<ComboboxSelected>>", update_teachers)
    existing_box.bind("<<ComboboxSelected>>", update_divisions)
    new_entry.bind("<KeyRelease>", update_info)
    amount_box.bind("<KeyRelease>", update_info)
    amount_box.bind("<<Increment>>", update_info)
    amount_box.bind("<<Decrement>>", update_info)

    footer = tk.Frame(outer, bg=BG)
    footer.pack(fill="x", pady=(16, 0))

    def run_split():
        grade = grade_var.get().strip()
        division = division_var.get().strip()
        subject = subject_var.get().strip()
        old_teacher = existing_var.get().strip()
        new_teacher = new_teacher_var.get().strip()
        try:
            count = int(transfer_var.get())
        except ValueError:
            messagebox.showerror("Invalid period count", "Enter a whole number of periods to transfer.", parent=window)
            return
        if not division:
            messagebox.showerror(
                "Division required", "Choose a division where the existing teacher teaches this subject.", parent=window
            )
            return
        matching = _teacher_subject_division_counts(source_workbook, grade, subject, old_teacher)
        if division not in matching:
            messagebox.showerror(
                "No matching periods", "The selected teacher does not teach this subject in that division.", parent=window
            )
            return
        if count < 1 or count > matching[division]:
            messagebox.showerror(
                "Invalid period count",
                f"{old_teacher} teaches {matching[division]} period(s) in Grade {grade}-{division}. "
                f"Enter a value from 1 to {matching[division]}.", parent=window,
            )
            return
        if not new_teacher:
            messagebox.showerror("New teacher required", "Enter the new teacher's name.", parent=window)
            return
        if new_teacher.casefold() == old_teacher.casefold():
            messagebox.showerror("Teacher names must differ", "Choose a different name for the new teacher.", parent=window)
            return
        subjects, _teachers, _counts = _split_subject_data(source_workbook, grade)
        if subject not in subjects:
            messagebox.showerror("Invalid subject", "Choose a subject listed for the selected grade.", parent=window)
            return
        confirmed = messagebox.askyesno(
            "Confirm subject split",
            f"Grade: {grade}\nDivision: {division}\nSubject: {subject}\nExisting teacher: {old_teacher}\nNew teacher: {new_teacher}\n\n"
            f"Transfer {count} period(s) in Grade {grade}-{division}.\n"
            "Only this division will be changed. Class weekly totals will remain unchanged. "
            "The new teacher must be available in every selected period.\n\nContinue?",
            parent=window,
        )
        if not confirmed:
            return
        try:
            result = split_subject_periods(
                workbook_path, grade, subject, old_teacher, new_teacher, count,
                division=division,
            )
        except Exception as error:
            messagebox.showerror("Could not split subject periods", str(error), parent=window)
            return
        try:
            if on_saved:
                on_saved(workbook_path)
        except Exception:
            pass
        source_workbook.close()
        messagebox.showinfo(
            "Subject periods split",
            f"Successfully assigned {result['periods_transferred']} {subject} period(s) to {result['new_teacher']} in "
            f"Grade {grade}-{division}.\n\n{result['existing_teacher']} keeps the remaining periods in this division.\n"
            "Class timetable totals, teacher sheets and Requirements counts have been updated.",
            parent=window,
        )
        window.destroy()

    tk.Button(footer, text="Cancel", command=window.destroy,
              bg="#F1F1EF", fg=TEXT, relief="flat", bd=0,
              padx=18, pady=10, font=("Segoe UI", 9, "bold"), cursor="hand2").pack(side="right")
    tk.Button(footer, text="Split periods", command=run_split,
              bg=PURPLE, fg=WHITE, activebackground=PURPLE_DARK,
              activeforeground=WHITE, relief="flat", bd=0,
              padx=20, pady=10, font=("Segoe UI", 10, "bold"), cursor="hand2").pack(side="right", padx=(0, 8))

    update_subjects()
    window.protocol("WM_DELETE_WINDOW", lambda: (source_workbook.close(), window.destroy()))
    window.transient(parent)
    window.grab_set()
    return window



# Create the workbook editor, optionally focused on a single class sheet.
def open_excel_editor(parent, workbook_path, on_saved=None, initial_sheet=None, locked_class=False, subject_first_mode=False):
    """Open a workbook editor, optionally focused on one class with subject-first editing."""
    if not os.path.isfile(workbook_path):
        messagebox.showerror("Workbook not found", "Generate a timetable or choose an existing .xlsx workbook first.", parent=parent)
        return None
    return ExcelTimetableEditor(
        parent, workbook_path, on_saved=on_saved, initial_sheet=initial_sheet,
        locked_class=locked_class, subject_first_mode=subject_first_mode,
    )

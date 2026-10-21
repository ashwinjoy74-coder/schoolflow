# Read a generated SchoolFlow workbook so timetables can be shown inside the app.
#
# Class sheets are named like "7-B" and teacher sheets like "T012". Each grid has
# a title in row 1, "Day" / "Period N" headings in row 2, and one row per day.
# Class cells hold "Subject\nTeacher"; teacher cells hold "Class\nSubject".

import os
import re

from openpyxl import load_workbook

CLASS_SHEET = re.compile(r"^(\d{1,2})-([A-Z])$", re.I)
TEACHER_SHEET = re.compile(r"^T\d{3,}$", re.I)


# Convert one worksheet's rows into title, day names, period names, and cell text.
def _read_grid(rows):
    title = str(rows[0][0] or "") if rows and rows[0] else ""
    header = rows[1] if len(rows) > 1 else ()
    periods = [str(value) for value in header[1:] if value not in (None, "")]
    days = []
    cells = []
    for row in rows[2:]:
        if not row or row[0] in (None, ""):
            continue
        days.append(str(row[0]))
        values = list(row[1:1 + len(periods)])
        values += [None] * (len(periods) - len(values))
        cells.append(["" if value is None else str(value) for value in values])
    return {"title": title, "days": days, "periods": periods, "cells": cells}


# Load every class and teacher timetable in the workbook.
def read_timetables(path):
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        # The Summary sheet lists each teacher's sheet name and subject in columns D–F.
        teacher_subjects = {}
        if "Summary" in workbook.sheetnames:
            for row in workbook["Summary"].iter_rows(min_row=2, min_col=4, max_col=6, values_only=True):
                if row and row[1]:
                    teacher_subjects[str(row[1])] = str(row[2] or "")

        classes = []
        teachers = []
        for name in workbook.sheetnames:
            match = CLASS_SHEET.match(name)
            is_teacher = bool(TEACHER_SHEET.match(name))
            if not match and not is_teacher:
                continue
            grid = _read_grid(list(workbook[name].iter_rows(values_only=True)))
            grid["sheet"] = name
            if match:
                grid["label"] = name
                grid["sort_key"] = (int(match.group(1)), match.group(2).upper())
                classes.append(grid)
            else:
                teacher = grid["title"].split(":", 1)[-1].strip() or name
                grid["label"] = teacher
                grid["subject"] = teacher_subjects.get(name, "")
                grid["sort_key"] = teacher.casefold()
                teachers.append(grid)
    finally:
        workbook.close()

    classes.sort(key=lambda grid: grid["sort_key"])
    teachers.sort(key=lambda grid: grid["sort_key"])
    return {
        "path": path,
        "modified": os.path.getmtime(path),
        "classes": classes,
        "teachers": teachers,
    }

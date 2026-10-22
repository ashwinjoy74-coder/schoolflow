# School Timetable Generator
# This program makes timetables for students and teachers.
# It stops a class or teacher from having two lessons at the same time.

import csv
import os
import sys
import threading

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

try:
    from ortools.sat.python import cp_model
except ImportError:
    cp_model = None


# ---------------- SCHOOL SETTINGS ----------------

DAYS = ["Monday", "Tuesday", "Wednesday", "Thursday", "Friday"]
PERIODS_PER_DAY = 7
CLASSES = range(1, 13)
DIVISIONS = ["A", "B", "C", "D"]

# The code lives in Program; input data and exports are kept outside it.
FOLDER = os.path.dirname(os.path.abspath(__file__))
PROJECT_FOLDER = os.path.dirname(FOLDER)
REQUIREMENTS_FOLDER = os.path.join(PROJECT_FOLDER, "Requirements")
EXPORTS_FOLDER = os.path.join(PROJECT_FOLDER, "Exports")
INPUT_FILE = os.path.join(REQUIREMENTS_FOLDER, "requirements.csv")
OUTPUT_FILE = os.path.join(EXPORTS_FOLDER, "school_timetable.xlsx")

# Longest time the solver may search before giving up, in seconds.
SOLVER_TIME_LIMIT = 180

# The functions below share the module settings above, so only one run may use them at a time.
_RUN_LOCK = threading.Lock()


# Holds a record of any subject counts that were changed to fill the selected week.
PERIOD_ADJUSTMENTS = []

# Group periods such as AMD: rows with a "group" value. All divisions in a group take
# it at the same time, split into activities (Music, Karate, ...) with their own teachers.
GROUPS = []
# The slot chosen for every group period, as (group, slot) pairs.
GROUP_SLOTS = []
# Requirement row numbers marked "double": their periods come in back-to-back pairs.
DOUBLE_ROWS = set()
# Second line of a class cell that holds a group period, e.g. "AMD\nGroup: Music, Karate".
GROUP_MARK = "Group:"
YES_VALUES = {"yes", "y", "true", "1", "x"}


# Divide extra periods proportionally and distribute rounding leftovers fairly.
def spread_periods(amount, weights):
    """Share extra periods fairly, using the existing subject periods as weights."""
    if amount <= 0:
        return [0 for weight in weights]

    total_weight = sum(weights)
    if total_weight <= 0:
        weights = [1 for weight in weights]
        total_weight = len(weights)

    exact_shares = [amount * weight / total_weight for weight in weights]
    shares = [int(value) for value in exact_shares]
    left = amount - sum(shares)

    # Give leftover periods to the subjects with the largest decimal remainder.
    order = sorted(
        range(len(weights)),
        key=lambda index: exact_shares[index] - shares[index],
        reverse=True,
    )
    for index in order[:left]:
        shares[index] += 1

    return shares


# Read and validate requirements.csv, then balance each division to the selected week capacity.
def read_data():
    global PERIOD_ADJUSTMENTS, GROUPS, DOUBLE_ROWS
    requirements = []
    lessons = []
    group_rows = []
    PERIOD_ADJUSTMENTS = []
    GROUPS = []
    DOUBLE_ROWS = set()

    if not os.path.exists(INPUT_FILE):
        raise ValueError("requirements.csv was not found. Put it in the Requirements folder.")

    with open(INPUT_FILE, "r", newline="", encoding="utf-8-sig") as file:
        reader = csv.DictReader(file)

        needed_columns = ["class", "division", "subject", "teacher", "periods_per_week"]
        if not reader.fieldnames or any(column not in reader.fieldnames for column in needed_columns):
            raise ValueError("The CSV must have these columns: class, division, subject, teacher, periods_per_week")

        for row_number, row in enumerate(reader, start=2):
            grade_text = (row.get("class") or "").strip()
            division = (row.get("division") or "").strip().upper()
            subject = (row.get("subject") or "").strip()
            teacher = (row.get("teacher") or "").strip()
            periods_text = (row.get("periods_per_week") or "").strip()

            # Ignore empty starter rows in the CSV template.
            if not subject and not teacher and not periods_text:
                continue

            if not grade_text or not division or not subject or not teacher or not periods_text:
                raise ValueError("Row " + str(row_number) + " is missing some information.")

            try:
                grade = int(grade_text)
                periods = int(periods_text)
            except ValueError:
                raise ValueError("Row " + str(row_number) + ": class and periods_per_week must be numbers.")

            if grade not in CLASSES:
                raise ValueError("Row " + str(row_number) + ": class must be between 1 and 12.")
            if division not in DIVISIONS:
                raise ValueError("Row " + str(row_number) + ": division must be A, B, C or D.")
            if periods < 1:
                raise ValueError("Row " + str(row_number) + ": periods_per_week must be at least 1.")

            group_name = (row.get("group") or "").strip()
            double = (row.get("double") or "").strip().casefold() in YES_VALUES
            if group_name:
                group_rows.append((grade, division, subject, teacher, periods, group_name, double, row_number))
                continue
            if double:
                DOUBLE_ROWS.add(len(requirements))
            requirements.append([grade, division, subject, teacher, periods])

    if not requirements:
        raise ValueError("The CSV has no subjects yet. Fill in the subject, teacher and weekly periods.")
    if (DOUBLE_ROWS or any(item[6] for item in group_rows)) and PERIODS_PER_DAY < 2:
        raise ValueError("'Two periods together' needs at least 2 periods per day.")

    # A teacher must be assigned to exactly one subject throughout the school.
    teacher_subjects = {}
    teacher_display_names = {}
    teacher_conflicts = {}
    for row in requirements:
        teacher_key = row[3].strip().casefold()
        subject = row[2].strip()
        teacher_display_names.setdefault(teacher_key, row[3].strip())
        previous_subject = teacher_subjects.get(teacher_key)
        if previous_subject and previous_subject.casefold() != subject.casefold():
            teacher_conflicts.setdefault(teacher_key, set()).update([previous_subject, subject])
        else:
            teacher_subjects[teacher_key] = subject

    if teacher_conflicts:
        messages = []
        for teacher_key, subjects in sorted(teacher_conflicts.items()):
            messages.append(teacher_display_names[teacher_key] + " is assigned to multiple subjects: " + ", ".join(sorted(subjects)))
        raise ValueError("Each teacher can teach only one subject. Correct these assignments in requirements.csv:\n- " + "\n- ".join(messages))

    # Standardize capitalization so one teacher always gets one workbook sheet.
    for row in requirements:
        row[3] = teacher_display_names[row[3].strip().casefold()]

    # Collect group rows into groups: one per grade and group name. Group teachers may
    # also teach a normal subject; the one-subject rule above covers normal rows only.
    groups_by_key = {}
    for grade, division, subject, teacher, periods, name, double, row_number in group_rows:
        teacher = teacher_display_names.setdefault(teacher.casefold(), teacher)
        group = groups_by_key.setdefault((grade, name.casefold()), {
            "name": name, "grade": grade, "divisions": [], "activities": [], "periods": periods, "double": False,
        })
        if periods != group["periods"]:
            raise ValueError("Row " + str(row_number) + ": every row of group " + group["name"] + " in class "
                             + str(grade) + " must have the same periods_per_week (" + str(group["periods"]) + ").")
        group["double"] = group["double"] or double
        if division not in group["divisions"]:
            group["divisions"].append(division)
        if (subject.casefold(), teacher.casefold()) not in [(s.casefold(), t.casefold()) for s, t in group["activities"]]:
            group["activities"].append((subject, teacher))
    for group in sorted(groups_by_key.values(), key=lambda item: (item["grade"], item["name"].casefold())):
        group["divisions"].sort(key=DIVISIONS.index)
        group["classes"] = [str(group["grade"]) + "-" + division for division in group["divisions"]]
        group["teachers"] = sorted({teacher for _subject, teacher in group["activities"]}, key=str.casefold)
        GROUPS.append(group)

    slots_per_week = len(DAYS) * PERIODS_PER_DAY

    # Group periods take fixed slots, so each division has fewer left for its subjects.
    group_periods_by_class = {}
    group_load = {}
    for group in GROUPS:
        for class_name in group["classes"]:
            group_periods_by_class[class_name] = group_periods_by_class.get(class_name, 0) + group["periods"]
        for teacher in group["teachers"]:
            group_load[teacher] = group_load.get(teacher, 0) + group["periods"]
    for class_name, count in sorted(group_periods_by_class.items()):
        if count > slots_per_week:
            raise ValueError("Class " + class_name + " has " + str(count) + " group periods, but the selected week has only "
                             + str(slots_per_week) + " periods.")
    for teacher, count in sorted(group_load.items()):
        if count > slots_per_week:
            raise ValueError(teacher + " is needed for " + str(count) + " group periods, but the selected week has only "
                             + str(slots_per_week) + " periods.")
    capacity_by_class = {
        str(grade) + "-" + division: slots_per_week - group_periods_by_class.get(str(grade) + "-" + division, 0)
        for grade in CLASSES for division in DIVISIONS
    }

    # Group input rows by division so that weekly periods can be adjusted together.
    rows_by_class = {}
    for row_index, row in enumerate(requirements):
        grade, division, subject, teacher, periods = row
        class_name = str(grade) + "-" + division
        if class_name not in rows_by_class:
            rows_by_class[class_name] = []
        rows_by_class[class_name].append(row_index)

    # Every division must have at least one subject row (unless groups fill its whole week).
    missing_classes = []
    for grade in CLASSES:
        for division in DIVISIONS:
            class_name = str(grade) + "-" + division
            if class_name not in rows_by_class and capacity_by_class[class_name] > 0:
                missing_classes.append(class_name)

    if missing_classes:
        raise ValueError("Add subject information for these divisions: " + ", ".join(missing_classes))

    # First choose a fair, per-division subject distribution for the selected week.
    # Short weeks can otherwise overload a teacher who teaches the same subject
    # in many divisions, so we rebalance repeated periods across subjects below.
    original_counts = [row[4] for row in requirements]
    adjusted_counts = original_counts[:]

    for class_name, row_indexes in rows_by_class.items():
        current_counts = [original_counts[index] for index in row_indexes]
        current_total = sum(current_counts)
        subject_count = len(row_indexes)
        capacity = capacity_by_class[class_name]

        if capacity < subject_count:
            group_note = (" after " + str(slots_per_week - capacity) + " group periods"
                          if capacity < slots_per_week else "")
            raise ValueError(
                class_name + " has " + str(subject_count) + " subject rows, but the selected week has only "
                + str(capacity) + " periods" + group_note + ". Each subject needs at least one period to be included. "
                + "Select more periods per day or more school days."
            )

        if current_total < capacity:
            extra_periods = capacity - current_total
            extra_by_subject = spread_periods(extra_periods, current_counts)
            new_counts = [current_counts[i] + extra_by_subject[i] for i in range(subject_count)]
        elif current_total > capacity:
            periods_to_share = capacity - subject_count
            weights = [max(count - 1, 0) for count in current_counts]
            extra_by_subject = spread_periods(periods_to_share, weights)
            new_counts = [1 + extra_by_subject[i] for i in range(subject_count)]
        else:
            new_counts = current_counts[:]

        for i, row_index in enumerate(row_indexes):
            adjusted_counts[row_index] = new_counts[i]

    # A teacher needs at least one slot for each class/division assignment.
    # If even this minimum is too large for the selected week, no clash-free
    # solution can exist without changing assignments or increasing capacity.
    teacher_minimum_load = dict(group_load)
    for row in requirements:
        teacher_minimum_load[row[3]] = teacher_minimum_load.get(row[3], 0) + 1

    impossible_teachers = [
        (teacher, count) for teacher, count in teacher_minimum_load.items()
        if count > slots_per_week
    ]
    if impossible_teachers:
        impossible_teachers.sort(key=lambda item: item[1], reverse=True)
        details = [
            teacher + " is assigned to " + str(count) + " class/division rows, but only "
            + str(slots_per_week) + " periods are available in the selected week."
            for teacher, count in impossible_teachers[:8]
        ]
        raise ValueError(
            "This timetable cannot fit into " + str(len(DAYS)) + " day(s) × "
            + str(PERIODS_PER_DAY) + " periods without teacher clashes. "
            + "Some teachers are assigned to more divisions than there are weekly time slots. "
            + "Select more periods per day/school days, or assign another teacher to some divisions.\n- "
            + "\n- ".join(details)
        )

    # Keep each class full, while moving extra repeats away from teachers whose
    # weekly workload exceeds the available slots. A moved period stays in the
    # same class and is assigned to another subject teacher who has free capacity.
    teacher_load = dict(group_load)
    for row_index, row in enumerate(requirements):
        teacher = row[3]
        teacher_load[teacher] = teacher_load.get(teacher, 0) + adjusted_counts[row_index]

    while True:
        overloaded = [teacher for teacher, total in teacher_load.items() if total > slots_per_week]
        if not overloaded:
            break

        overloaded_teacher = max(overloaded, key=lambda teacher: teacher_load[teacher])
        candidates = []
        for class_name, row_indexes in rows_by_class.items():
            donor_rows = [
                index for index in row_indexes
                if requirements[index][3] == overloaded_teacher and adjusted_counts[index] > 1
            ]
            if not donor_rows:
                continue

            for donor_index in donor_rows:
                for receiver_index in row_indexes:
                    if receiver_index == donor_index:
                        continue
                    receiver_teacher = requirements[receiver_index][3]
                    if receiver_teacher == overloaded_teacher:
                        continue
                    free_slots = slots_per_week - teacher_load.get(receiver_teacher, 0)
                    if free_slots <= 0:
                        continue

                    # Prefer giving the repeat to a subject that is below its
                    # original period target, then to the teacher with more room.
                    target = max(original_counts[receiver_index], 1)
                    ratio_after_move = (adjusted_counts[receiver_index] + 1) / target
                    candidates.append((
                        ratio_after_move,
                        -free_slots,
                        adjusted_counts[receiver_index],
                        class_name,
                        donor_index,
                        receiver_index,
                    ))

        if not candidates:
            raise ValueError(
                "The timetable needs " + str(slots_per_week) + " periods per division, but the current teacher "
                + "assignments cannot fit without a clash. Some teachers are already needed in too many divisions. "
                + "Add another teacher for the subject or choose more school days/periods per day."
            )

        candidates.sort()
        _, _, _, _, donor_index, receiver_index = candidates[0]
        donor_teacher = requirements[donor_index][3]
        receiver_teacher = requirements[receiver_index][3]
        adjusted_counts[donor_index] -= 1
        adjusted_counts[receiver_index] += 1
        teacher_load[donor_teacher] -= 1
        teacher_load[receiver_teacher] = teacher_load.get(receiver_teacher, 0) + 1

    # Store the final adjusted counts and record the changes for Excel reporting.
    for row_index, new_count in enumerate(adjusted_counts):
        grade, division, subject, teacher, old_count = requirements[row_index]
        if old_count != new_count:
            PERIOD_ADJUSTMENTS.append({
                "class_name": str(grade) + "-" + division,
                "grade": grade,
                "division": division,
                "subject": subject,
                "teacher": teacher,
                "original_periods": old_count,
                "adjusted_periods": new_count,
            })
        requirements[row_index][4] = new_count

    # Create one lesson entry for every adjusted weekly period. In a "double" row every
    # second lesson "follows" the one before it, i.e. comes straight after it.
    lessons = []
    for row_index, (grade, division, subject, teacher, periods) in enumerate(requirements):
        class_name = str(grade) + "-" + division
        for number in range(periods):
            lesson = {
                "class_name": class_name,
                "subject": subject,
                "teacher": teacher,
            }
            if row_index in DOUBLE_ROWS and number % 2 == 1:
                lesson["follows"] = len(lessons) - 1
            lessons.append(lesson)

    # Final validation: every teacher must fit into the selected week's slots.
    final_teacher_load = dict(group_load)
    for lesson in lessons:
        teacher = lesson["teacher"]
        final_teacher_load[teacher] = final_teacher_load.get(teacher, 0) + 1

    teacher_problems = [
        teacher + " has " + str(total) + " lessons, but only " + str(slots_per_week) + " periods are available."
        for teacher, total in final_teacher_load.items() if total > slots_per_week
    ]
    if teacher_problems:
        raise ValueError("Teacher workloads exceed the weekly capacity after adjusting subject repeats:\n- "
                         + "\n- ".join(teacher_problems))

    return requirements, lessons


# Find a period for every lesson without creating class or teacher clashes.
# Use the constraint solver to assign every lesson without class or teacher clashes.
def make_timetable(lessons):
    if cp_model is None:
        raise RuntimeError("OR-Tools is not installed. Install it with: python -m pip install -r requirements.txt")

    global GROUP_SLOTS
    total_slots = len(DAYS) * PERIODS_PER_DAY
    model = cp_model.CpModel()
    choices = []

    # A lesson gets one choice for the day and period.
    for lesson_number in range(len(lessons)):
        lesson_choices = []
        for slot in range(total_slots):
            lesson_choices.append(model.NewBoolVar("lesson_" + str(lesson_number) + "_slot_" + str(slot)))
        model.AddExactlyOne(lesson_choices)
        choices.append(lesson_choices)

    # Each group period is one choice shared by all of the group's divisions and teachers.
    sessions = [(group, number) for group in GROUPS for number in range(group["periods"])]
    session_choices = []
    for session_number in range(len(sessions)):
        options = [model.NewBoolVar("group_" + str(session_number) + "_slot_" + str(slot)) for slot in range(total_slots)]
        model.AddExactlyOne(options)
        session_choices.append(options)

    # Collect the choices that use each class and each teacher.
    class_lessons = {}
    teacher_lessons = {}

    for lesson_number, lesson in enumerate(lessons):
        class_lessons.setdefault(lesson["class_name"], []).append(choices[lesson_number])
        teacher_lessons.setdefault(lesson["teacher"], []).append(choices[lesson_number])
    for session_number, (group, _number) in enumerate(sessions):
        for class_name in group["classes"]:
            class_lessons.setdefault(class_name, []).append(session_choices[session_number])
        for teacher in group["teachers"]:
            teacher_lessons.setdefault(teacher, []).append(session_choices[session_number])

    # Every class must have exactly one lesson in every period.
    # This prevents free periods in student/class timetables.
    for class_name in class_lessons:
        for slot in range(total_slots):
            model.AddExactlyOne(options[slot] for options in class_lessons[class_name])

    # A teacher can only teach one class in a period.
    for teacher in teacher_lessons:
        for slot in range(total_slots):
            model.AddAtMostOne(options[slot] for options in teacher_lessons[teacher])

    # "Two periods together": the second of a pair sits in the next period of the same day.
    def keep_together(first, second):
        for slot in range(total_slots):
            if slot % PERIODS_PER_DAY == PERIODS_PER_DAY - 1:
                model.Add(first[slot] == 0)
            else:
                model.Add(second[slot + 1] == first[slot])
            if slot % PERIODS_PER_DAY == 0:
                model.Add(second[slot] == 0)

    for lesson_number, lesson in enumerate(lessons):
        if "follows" in lesson:
            keep_together(choices[lesson["follows"]], choices[lesson_number])
    for session_number, (group, number) in enumerate(sessions):
        if group["double"] and number % 2 == 1:
            keep_together(session_choices[session_number - 1], session_choices[session_number])

    solver = cp_model.CpSolver()
    solver.parameters.max_time_in_seconds = SOLVER_TIME_LIMIT
    solver.parameters.num_search_workers = 8
    result = solver.Solve(model)

    if result != cp_model.OPTIMAL and result != cp_model.FEASIBLE:
        raise RuntimeError("A timetable could not be made with these requirements. Check the teacher assignments, "
                           "weekly periods, groups and 'two periods together' settings.")

    def chosen(options):
        return next(slot for slot in range(total_slots) if solver.Value(options[slot]) == 1)

    # Save the chosen period number for every lesson and group period.
    chosen_slots = [chosen(options) for options in choices]
    GROUP_SLOTS = [(group, chosen(session_choices[number])) for number, (group, _n) in enumerate(sessions)]
    return chosen_slots


# Make the Excel workbook with one sheet for every class division and teacher.
# Create the summary, requirements, class timetable, and teacher timetable worksheets.
def save_excel(requirements, lessons, chosen_slots):
    class_names = []
    teacher_names = []

    for grade in CLASSES:
        for division in DIVISIONS:
            class_names.append(str(grade) + "-" + division)

    teacher_subject_map = {}
    for lesson in lessons:
        if lesson["teacher"] not in teacher_names:
            teacher_names.append(lesson["teacher"])
        existing_subject = teacher_subject_map.get(lesson["teacher"])
        if existing_subject and existing_subject.casefold() != lesson["subject"].casefold():
            raise RuntimeError(
                lesson["teacher"] + " is assigned to more than one subject: "
                + existing_subject + " and " + lesson["subject"] + "."
            )
        teacher_subject_map[lesson["teacher"]] = lesson["subject"]
    # Teachers who only take group activities are listed as e.g. "AMD: Music".
    for group in GROUPS:
        for subject, teacher in group["activities"]:
            if teacher not in teacher_names:
                teacher_names.append(teacher)
            teacher_subject_map.setdefault(teacher, group["name"] + ": " + subject)
    teacher_names.sort()

    # Timetable sizes use the selected number of days and periods.
    class_tables = {}
    teacher_tables = {}

    for class_name in class_names:
        class_tables[class_name] = [["" for period in range(PERIODS_PER_DAY)] for day in DAYS]

    for teacher in teacher_names:
        teacher_tables[teacher] = [["" for period in range(PERIODS_PER_DAY)] for day in DAYS]

    # Fill the timetable tables using the periods found by the solver.
    for lesson_number, lesson in enumerate(lessons):
        slot = chosen_slots[lesson_number]
        day_number = slot // PERIODS_PER_DAY
        period_number = slot % PERIODS_PER_DAY
        class_name = lesson["class_name"]
        teacher = lesson["teacher"]

        # Check again so clashes are not accidentally written to Excel.
        if class_tables[class_name][day_number][period_number] != "":
            raise RuntimeError("A class clash was found while creating the Excel file.")
        if teacher_tables[teacher][day_number][period_number] != "":
            raise RuntimeError("A teacher clash was found while creating the Excel file.")

        class_tables[class_name][day_number][period_number] = lesson["subject"] + "\n" + teacher
        teacher_tables[teacher][day_number][period_number] = class_name + "\n" + lesson["subject"]

    # Group periods: every division shows the group and its activities; every activity
    # teacher shows the divisions and the activity they take.
    for group, slot in GROUP_SLOTS:
        day_number = slot // PERIODS_PER_DAY
        period_number = slot % PERIODS_PER_DAY
        class_text = group["name"] + "\n" + GROUP_MARK + " " + ", ".join(subject for subject, _t in group["activities"])
        for class_name in group["classes"]:
            if class_tables[class_name][day_number][period_number] != "":
                raise RuntimeError("A class clash was found while creating the Excel file.")
            class_tables[class_name][day_number][period_number] = class_text
        for teacher in group["teachers"]:
            if teacher_tables[teacher][day_number][period_number] != "":
                raise RuntimeError("A teacher clash was found while creating the Excel file.")
            subjects = " / ".join(subject for subject, name in group["activities"] if name == teacher)
            teacher_tables[teacher][day_number][period_number] = (", ".join(group["classes"]) + "\n"
                                                                  + group["name"] + ": " + subjects)

    # Double-check that every class period has been filled before saving.
    for class_name in class_names:
        for day_number, day in enumerate(DAYS):
            for period_number in range(PERIODS_PER_DAY):
                if class_tables[class_name][day_number][period_number] == "":
                    raise RuntimeError(
                        "No-free-period check failed for Class " + class_name +
                        " on " + day + ", Period " + str(period_number + 1) + "."
                    )

    workbook = Workbook()
    summary = workbook.active
    summary.title = "Summary"
    summary["A1"] = "School Timetable"
    summary["A1"].font = Font(size=16, bold=True, color="FFFFFF")
    summary["A1"].fill = PatternFill("solid", fgColor="244062")
    summary.append(["School days", len(DAYS)])
    summary.append(["Periods per day", PERIODS_PER_DAY])
    summary.append(["Number of class divisions", len(class_names)])
    summary.append(["Number of teachers", len(teacher_names)])
    summary.append(["Lessons scheduled", len(lessons)])
    summary.append(["Clash check", "No class or teacher is double-booked."])
    summary.append(["Class free periods", 0])
    summary.append(["Class timetable check", "Every division has a lesson in every selected period."])
    summary.append(["Teacher free periods", "May exist when a teacher is not assigned to a class."])
    summary.append(["Teacher subject rule", "Each teacher is assigned to one subject across all grades and divisions."])
    summary.append(["Teacher workbook sheets", "Each teacher has one combined timetable sheet listing all classes they teach."])
    summary.append(["Automatic subject repeats", "Adjusted to fill the selected weekly capacity where needed."])
    summary.append(["Changed subject rows", len(PERIOD_ADJUSTMENTS)])
    summary.append(["Group periods", sum(group["periods"] for group in GROUPS)])
    summary.column_dimensions["A"].width = 28
    summary.column_dimensions["B"].width = 65

    # Include the input data in the workbook too.
    data_sheet = workbook.create_sheet("Requirements")
    data_sheet.append(["class", "division", "subject", "teacher", "periods_per_week", "double"])
    for row_index, row in enumerate(requirements):
        data_sheet.append(row + ["yes" if row_index in DOUBLE_ROWS else ""])

    for cell in data_sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="244062")

    for column, width in zip(["A", "B", "C", "D", "E", "F"], [12, 12, 24, 28, 20, 10]):
        data_sheet.column_dimensions[column].width = width

    # Group periods are listed on their own sheet, one row per activity.
    if GROUPS:
        group_sheet = workbook.create_sheet("Groups")
        group_sheet.append(["Group", "Class", "Divisions", "Activity", "Teacher", "Periods per week", "Two periods together"])
        for group in GROUPS:
            for subject, teacher in group["activities"]:
                group_sheet.append([group["name"], group["grade"], ", ".join(group["divisions"]), subject, teacher,
                                    group["periods"], "yes" if group["double"] else ""])
        for cell in group_sheet[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="244062")
        for column, width in zip("ABCDEFG", [16, 8, 14, 22, 28, 16, 20]):
            group_sheet.column_dimensions[column].width = width

    # Keep a clear record of any automatic repeat changes made for this timetable.
    if PERIOD_ADJUSTMENTS:
        adjustment_sheet = workbook.create_sheet("Repeat Adjustments")
        adjustment_sheet.append(["Class", "Subject", "Teacher", "Original periods", "Scheduled periods"])
        for item in PERIOD_ADJUSTMENTS:
            adjustment_sheet.append([
                item["class_name"], item["subject"], item["teacher"],
                item["original_periods"], item["adjusted_periods"]
            ])
        for cell in adjustment_sheet[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="244062")
        for column, width in zip(["A", "B", "C", "D", "E"], [14, 24, 30, 20, 20]):
            adjustment_sheet.column_dimensions[column].width = width

    # Add one sheet per class division.
    for class_name in class_names:
        sheet = workbook.create_sheet(class_name)
        add_timetable_to_sheet(sheet, "Class " + class_name, class_tables[class_name])

    # Excel sheet names have a 31-character limit, so use T001, T002, etc.
    # The teacher's full name is shown as the heading inside the sheet.
    teacher_sheet_names = {}
    for number, teacher in enumerate(teacher_names, start=1):
        sheet_name = "T" + str(number).zfill(3)
        teacher_sheet_names[teacher] = sheet_name
        sheet = workbook.create_sheet(sheet_name)
        add_timetable_to_sheet(sheet, "Teacher: " + teacher, teacher_tables[teacher])

    # Add a list of teacher names and their worksheet names to the summary.
    summary["D1"] = "Teacher"
    summary["E1"] = "Sheet"
    summary["F1"] = "Subject"
    for header_cell in (summary["D1"], summary["E1"], summary["F1"]):
        header_cell.font = Font(bold=True, color="FFFFFF")
        header_cell.fill = PatternFill("solid", fgColor="244062")
    for row_number, teacher in enumerate(teacher_names, start=2):
        summary.cell(row=row_number, column=4, value=teacher)
        summary.cell(row=row_number, column=5, value=teacher_sheet_names[teacher])
        summary.cell(row=row_number, column=6, value=teacher_subject_map[teacher])
    summary.column_dimensions["D"].width = 32
    summary.column_dimensions["E"].width = 12
    summary.column_dimensions["F"].width = 24

    output_folder = os.path.dirname(os.path.abspath(OUTPUT_FILE))
    os.makedirs(output_folder, exist_ok=True)
    workbook.save(OUTPUT_FILE)


# Format the timetable so it is easy to read in Excel.
# Format one weekly timetable grid consistently in an Excel worksheet.
def add_timetable_to_sheet(sheet, title, timetable):
    sheet.merge_cells(start_row=1, start_column=1, end_row=1, end_column=PERIODS_PER_DAY + 1)
    sheet.cell(row=1, column=1, value=title)
    sheet.cell(row=1, column=1).font = Font(size=15, bold=True, color="FFFFFF")
    sheet.cell(row=1, column=1).fill = PatternFill("solid", fgColor="244062")
    sheet.cell(row=1, column=1).alignment = Alignment(horizontal="center")

    sheet.cell(row=2, column=1, value="Day")
    for period in range(PERIODS_PER_DAY):
        sheet.cell(row=2, column=period + 2, value="Period " + str(period + 1))

    for cell in sheet[2]:
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="D9EAF7")
        cell.alignment = Alignment(horizontal="center", vertical="center")

    for day_number, day in enumerate(DAYS):
        sheet.cell(row=day_number + 3, column=1, value=day)
        sheet.cell(row=day_number + 3, column=1).font = Font(bold=True)
        for period in range(PERIODS_PER_DAY):
            cell = sheet.cell(row=day_number + 3, column=period + 2, value=timetable[day_number][period])
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    sheet.column_dimensions["A"].width = 15
    for column in range(2, PERIODS_PER_DAY + 2):
        sheet.column_dimensions[chr(64 + column)].width = 18
    for row in range(3, len(DAYS) + 3):
        sheet.row_dimensions[row].height = 48

    sheet.freeze_panes = "B3"
    sheet.sheet_view.showGridLines = False


# Apply one run's settings to the module before calling the steps above.
def _configure(input_file, days, periods_per_day, output_file=None, time_limit=None):
    global INPUT_FILE, OUTPUT_FILE, DAYS, PERIODS_PER_DAY, SOLVER_TIME_LIMIT
    INPUT_FILE = input_file
    DAYS = list(days)
    PERIODS_PER_DAY = int(periods_per_day)
    if output_file:
        OUTPUT_FILE = output_file
    if time_limit:
        SOLVER_TIME_LIMIT = int(time_limit)


# Check requirements for a week without solving it; raises ValueError with the reason.
def check_requirements(input_file, days, periods_per_day):
    with _RUN_LOCK:
        _configure(input_file, days, periods_per_day)
        requirements, lessons = read_data()
        return {
            "rows": len(requirements),
            "lessons": len(lessons),
            "teachers": len(set(lesson["teacher"] for lesson in lessons) | set(t for g in GROUPS for t in g["teachers"])),
            "adjusted_divisions": len(set(item["class_name"] for item in PERIOD_ADJUSTMENTS)),
            "groups": len(GROUPS),
            "group_periods": sum(group["periods"] for group in GROUPS),
        }


# Run the full pipeline for the app: validate, solve, and save the workbook.
# `log` is called with short progress messages from the worker thread.
def generate(input_file, output_file, days, periods_per_day, time_limit=None, log=None):
    log = log or (lambda message: None)
    with _RUN_LOCK:
        _configure(input_file, days, periods_per_day, output_file, time_limit)
        requirements, lessons = read_data()
        adjusted_classes = sorted(set(item["class_name"] for item in PERIOD_ADJUSTMENTS))
        if adjusted_classes:
            log("Adjusted subject repeats for " + str(len(adjusted_classes)) + " divisions to fill "
                + str(len(DAYS) * PERIODS_PER_DAY) + " weekly periods each.")
        group_note = (" and " + str(sum(group["periods"] for group in GROUPS)) + " group periods" if GROUPS else "")
        log("Input checked. Scheduling " + str(len(lessons)) + " lessons" + group_note + " (up to "
            + str(SOLVER_TIME_LIMIT) + " s)…")
        chosen_slots = make_timetable(lessons)
        log("Schedule found. Writing the Excel workbook…")
        save_excel(requirements, lessons, chosen_slots)
        return {
            "output": OUTPUT_FILE,
            "lessons": len(lessons),
            "groups": len(GROUPS),
            "teachers": len(set(lesson["teacher"] for lesson in lessons) | set(t for g in GROUPS for t in g["teachers"])),
            "divisions": len(CLASSES) * len(DIVISIONS),
            "adjusted_divisions": len(adjusted_classes),
            "weekly_periods": len(DAYS) * PERIODS_PER_DAY,
        }


# Run validation, solve the schedule, save Excel, and report any expected errors.
def main():
    try:
        requirements, lessons = read_data()
        chosen_slots = make_timetable(lessons)
        save_excel(requirements, lessons, chosen_slots)
        print("Timetable created successfully!")
        print("Saved to:", OUTPUT_FILE)
        print("Class divisions:", len(CLASSES) * len(DIVISIONS))
        print("Teachers:", len(set(lesson["teacher"] for lesson in lessons)))
        print("Lessons scheduled:", len(lessons))
    except (ValueError, RuntimeError, OSError) as error:
        print("ERROR:", error)
        sys.exit(1)


if __name__ == "__main__":
    main()

# Lesson editing and period splitting for the app, without any editor window.
#
# The rules live in school_timetable_excel_editor: a lesson change updates the
# class sheet, the teacher's sheet and the Requirements count together, and is
# refused when it would double-book a teacher or give them a second subject.
# This module runs those same methods directly on the workbook.

import os
import threading

from openpyxl import load_workbook

from school_timetable_excel_editor import (
    CLASS_RE, ExcelTimetableEditor, _split_subject_data, _teacher_subject_division_counts, group_period_message,
    split_subject_periods,
)

_LOCK = threading.Lock()


class _Status:
    """Stands in for the editor window's status text, which is not shown here."""

    def set(self, _text):
        pass


class _WorkbookEditor(ExcelTimetableEditor):
    """The editor's lesson rules applied to a workbook file, with no window."""

    def __init__(self, path):
        self.path = os.path.abspath(path)
        self.workbook = load_workbook(self.path)
        self.teacher_sheets = {}
        self.dirty = False
        self.selected_cell = None
        self.status_var = _Status()
        self._refresh_teacher_sheet_map()

    def load_sheet(self):
        # The window version redraws its grid here; nothing to redraw.
        pass

    # Same checks the editor runs before saving: each teacher keeps one subject.
    def validate(self):
        seen = {}

        def check(teacher, subject):
            if not teacher or not subject:
                return
            key = teacher.strip().casefold()
            previous = seen.get(key)
            if previous and previous.casefold() != subject.strip().casefold():
                raise ValueError(teacher + " would teach both " + previous + " and " + subject
                                 + ". Each teacher can teach only one subject.")
            seen[key] = subject.strip()

        if "Requirements" in self.workbook.sheetnames:
            sheet = self.workbook["Requirements"]
            for row in range(2, sheet.max_row + 1):
                check(self._display_value(sheet.cell(row, 4).value).strip(),
                      self._display_value(sheet.cell(row, 3).value).strip())
        for name in self.workbook.sheetnames:
            if self._is_class_sheet(name):
                sheet = self.workbook[name]
                for row in range(3, sheet.max_row + 1):
                    for col in range(2, sheet.max_column + 1):
                        subject, teacher = self._split_lesson(self._display_value(sheet.cell(row, col).value))
                        check(teacher, subject)

    def save(self):
        self.validate()
        try:
            self.workbook.save(self.path)
        except PermissionError:
            raise ValueError("The workbook is open in Excel. Close it there, then try again.")


# Keep one loaded workbook between edits; reload when the file changes on disk.
_cache = {"path": None, "mtime": None, "editor": None}


def _editor(path):
    path = os.path.abspath(path)
    mtime = os.path.getmtime(path)
    if _cache["path"] != path or _cache["mtime"] != mtime or _cache["editor"] is None:
        _cache.update(path=path, mtime=mtime, editor=_WorkbookEditor(path))
    return _cache["editor"]


def _forget():
    _cache.update(path=None, mtime=None, editor=None)


# Subjects available to a class slot, each with its teachers and whether they are free then.
def lesson_options(path, class_name, day_index, period_index):
    with _LOCK:
        editor = _editor(path)
        if class_name not in editor.workbook.sheetnames:
            raise ValueError("Class " + class_name + " is not in this workbook.")
        row, col = int(day_index) + 3, int(period_index) + 2
        sheet = editor.workbook[class_name]
        if editor._is_group_cell(sheet.cell(row, col).value):
            raise ValueError(group_period_message(editor._display_value(sheet.cell(row, col).value)))
        subject, teacher = editor._split_lesson(editor._display_value(sheet.cell(row, col).value))
        by_subject = {}
        for pair in editor._teacher_subject_pairs_for_grade(class_name):
            busy = None
            if not (pair["teacher"].casefold() == teacher.casefold() and pair["subject"].casefold() == subject.casefold()):
                clash = editor._teacher_clash(pair["teacher"], class_name, row, col)
                if clash:
                    busy = clash[0]
            by_subject.setdefault(pair["subject"], []).append({"name": pair["teacher"], "busy": busy})
        return {
            "class": class_name,
            "day": editor._display_value(sheet.cell(row, 1).value),
            "period": editor._display_value(sheet.cell(2, col).value),
            "current": {"subject": subject, "teacher": teacher},
            "subjects": [{"subject": name, "teachers": teachers}
                         for name, teachers in sorted(by_subject.items(), key=lambda item: item[0].casefold())],
        }


# Put a lesson in a class slot (or clear it when subject/teacher are empty) and save.
def set_lesson(path, class_name, day_index, period_index, subject, teacher):
    with _LOCK:
        editor = _editor(path)
        row, col = int(day_index) + 3, int(period_index) + 2
        try:
            editor._apply_lesson_change(class_name, row, col, subject or "", teacher or "")
            editor.save()
        except Exception:
            _forget()  # drop the half-applied in-memory change
            raise
        _cache["mtime"] = os.path.getmtime(editor.path)
        return {"ok": True}


# Choices for the Split subject periods dialog.
def split_options(path, grade):
    with _LOCK:
        workbook = load_workbook(path, read_only=False)
        try:
            subjects, teachers_by_subject, _counts = _split_subject_data(workbook, grade)
            result = []
            for subject in subjects:
                teachers = []
                for teacher in sorted(teachers_by_subject.get(subject.casefold(), {}).values(), key=str.casefold):
                    divisions = _teacher_subject_division_counts(workbook, grade, subject, teacher)
                    if divisions:
                        teachers.append({"name": teacher, "divisions": divisions})
                if teachers:
                    result.append({"subject": subject, "teachers": teachers})
            return result
        finally:
            workbook.close()


# Teachers who already teach `subject` and could take over some of `existing_teacher`'s
# periods in the chosen grade/division. For each one, `free` is how many periods per
# division they can take without a clash (the smallest over the affected divisions).
def split_candidates(path, grade, subject, existing_teacher, division):
    grade = str(grade).strip()
    subject_key = subject.strip().casefold()
    existing_key = existing_teacher.strip().casefold()
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        names = {}      # teacher key -> display name (teaches this subject somewhere)
        lessons = {}    # teacher key -> lessons scheduled this week
        busy = {}       # (teacher key, row, col) -> class they teach then
        to_move = {}    # class name -> [(row, col)] slots the existing teacher holds
        if "Requirements" in workbook.sheetnames:
            for row in workbook["Requirements"].iter_rows(min_row=2, values_only=True):
                if row and len(row) >= 4 and row[2] and row[3] and str(row[2]).strip().casefold() == subject_key:
                    names.setdefault(str(row[3]).strip().casefold(), str(row[3]).strip())
        for name in workbook.sheetnames:
            if not CLASS_RE.match(name):
                continue
            for r, values in enumerate(workbook[name].iter_rows(min_row=3, values_only=True), start=3):
                for c, value in enumerate(values[1:], start=2):
                    cell_subject, cell_teacher = ExcelTimetableEditor._split_lesson(ExcelTimetableEditor._display_value(value))
                    if not cell_teacher:
                        continue
                    key = cell_teacher.casefold()
                    busy[(key, r, c)] = name
                    lessons[key] = lessons.get(key, 0) + 1
                    if cell_subject.casefold() == subject_key:
                        names.setdefault(key, cell_teacher)
                        target = name.split("-", 1)[0] == grade and (
                            not division or division == "All divisions" or name.split("-", 1)[1].casefold() == division.casefold())
                        if target and key == existing_key:
                            to_move.setdefault(name, []).append((r, c))
    finally:
        workbook.close()

    candidates = []
    for key, display in names.items():
        if key == existing_key:
            continue
        free_per_class = [sum(1 for r, c in slots if (key, r, c) not in busy) for slots in to_move.values()]
        candidates.append({"name": display, "lessons": lessons.get(key, 0),
                           "free": min(free_per_class) if free_per_class else 0})
    candidates.sort(key=lambda item: (-item["free"], item["lessons"], item["name"].casefold()))
    return candidates


def split(path, grade, subject, existing_teacher, new_teacher, periods, division):
    with _LOCK:
        _forget()
        try:
            return split_subject_periods(path, grade, subject, existing_teacher, new_teacher, periods,
                                         division=division or "All divisions")
        except PermissionError:
            raise ValueError("The workbook is open in Excel. Close it there, then try again.")
